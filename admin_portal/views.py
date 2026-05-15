import functools
import json
from datetime import date, timedelta
from decimal import Decimal

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db import IntegrityError, transaction
from django.db.models import DecimalField, F, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from accounts.models import (
    AccountTypeConfiguration,
    AuditLog,
    FeeCharge,
    FeeSchedule,
    FinancialPeriod,
    FundAccount,
    FundAllocationRule,
    FundTransaction,
    Guarantor,
    InterestReceivable,
    LoanAccount,
    LoanApplication,
    LoanRepayment,
    LoanTypeConfiguration,
    MemberAccount,
    ProfitAndLoss,
    ShareCapital,
    SocietyAccount,
    SocietyConfiguration,
    Transaction,
    User,
    Voucher,
    VoucherEntry,
)
from accounts.services import interest_engine
from accounts.services import loans as loan_service
from accounts.services import transactions as transaction_service
from accounts.services.exceptions import ServiceError
from accounts.utils import (
    LoanListItem,
    _get_client_ip,
    active_financial_period,
    apply_fund_allocations,
    ensure_member_submodels,
    financial_period_overlapping,
    format_current_address_display,
    format_permanent_address_display,
    get_account_type_cashflow,
    get_financial_summary,
    log_action,
    nominee_primary,
    nominee_secondary,
    persist_financial_snapshots,
    split_full_name,
    sync_loan_interest_receivables,
    user_active_share_capital_total,
    user_active_share_count,
    user_kyc,
    user_primary_share_lot,
    validate_password_strength,
)
from admin_portal.portal_fy import (
    PORTAL_FY_SESSION_KEY,
    effective_fy_window,
    report_dates_from_request,
    resolve_portal_financial_period,
    set_working_financial_period,
)

# Rate limiting cache for member creation (simple in-memory)
_member_creation_timestamps = {}


def _safe_same_origin_redirect_path(raw):
    """Allow only relative same-site paths (avoid open redirects)."""
    if not raw or not isinstance(raw, str):
        return None
    s = raw.strip()
    if not s.startswith("/") or s.startswith("//"):
        return None
    return s


def admin_required(view_func):
    """Decorator to check if user is staff"""

    @functools.wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_staff:
            return HttpResponseForbidden("Access denied. Administrators only.")
        return view_func(request, *args, **kwargs)

    return wrapper


def _credit_transaction(transaction_type):
    """Thin alias for `transaction_service.is_credit_transaction_type` (kept
    for in-file readability)."""
    return transaction_service.is_credit_transaction_type(transaction_type)


def _next_transaction_number():
    """Alias kept for any leftover callers in this file."""
    return transaction_service.next_transaction_number()


def _next_voucher_number():
    """Alias kept for any leftover callers in this file."""
    return transaction_service.next_voucher_number()


def _has_unpaid_emi(user):
    """Thin alias for `accounts.services.loans.has_unpaid_emi` (kept for
    backwards-compat with the two callers in this file)."""
    return loan_service.has_unpaid_emi(user)


def _default_account_rates():
    return {
        "fd": Decimal("7.50"),
        "cd": Decimal("7.00"),
        "rd": Decimal("7.00"),
        "od": Decimal("12.00"),
        "share": Decimal("0.00"),
        "sukanya": Decimal("8.00"),
        "suputra": Decimal("8.50"),
    }


def _process_transaction_line(user, line, payment_mode, reference_number, remarks, created_by, instrument=None):
    """Thin wrapper kept for any leftover callers — delegates to
    `transaction_service.post_transaction`.
    """
    return transaction_service.post_transaction(
        member_account_id=line["account_id"],
        transaction_type=line["transaction_type"],
        amount=line.get("amount", 0),
        description=line.get("description") or "",
        payment_mode=payment_mode,
        reference_number=reference_number or "",
        remarks=remarks or "",
        loan_repayment_id=line.get("loan_repayment_id") or None,
        instrument=instrument,
        instrument_payload=None,
        actor=created_by,
        audit=False,
        audit_via="panel",
    )


def _apply_member_status_rules(user, new_status, previous_status=None):
    today = timezone.now().date()
    previous_status = previous_status or user.status
    if new_status == "inactive":
        if _has_unpaid_emi(user):
            raise ValueError("Cannot mark member inactive while pending EMI payments exist.")
        if not user.inactive_since:
            user.inactive_since = today
    elif new_status == "active":
        user.inactive_since = None
    elif new_status == "resign":
        if previous_status != "inactive":
            raise ValueError("Member must be inactive before resignation.")
        if not user.inactive_since:
            raise ValueError("Inactive date is required before resignation.")
        if (today - user.inactive_since).days < 30:
            raise ValueError("Member can resign only after 30 days in inactive status.")
        if _has_unpaid_emi(user):
            raise ValueError("Cannot resign member while pending EMI payments exist.")
        user.exit_date = today
        user.is_active = False


@login_required
@admin_required
def home_view(request):
    """Admin dashboard view with stats, charts, and recent transactions"""
    now = timezone.now()
    today = now.date()

    # Member stats (exclude admin/staff accounts)
    total_members = User.objects.filter(is_deleted=False, is_superuser=False, role="member").count()
    active_members = User.objects.filter(is_deleted=False, is_superuser=False, role="member", status="active").count()

    # Account stats
    total_accounts = MemberAccount.objects.filter(is_deleted=False).count()
    active_accounts = MemberAccount.objects.filter(is_deleted=False, status="active").count()
    total_deposits = (
        MemberAccount.objects.filter(is_deleted=False, status="active").aggregate(total=Sum("balance"))["total"] or 0
    )

    # Account type distribution (for pie chart) - show total balance by type
    account_type_data = []
    for item in (
        MemberAccount.objects.filter(is_deleted=False, status="active")
        .values("account_type")
        .annotate(total=Sum("balance"))
        .order_by("account_type")
    ):
        account_type_data.append(
            {"account_type": item["account_type"], "total": float(item["total"]) if item["total"] else 0.0}
        )

    # Transaction stats
    total_transactions = Transaction.objects.count()
    today_transactions = Transaction.objects.filter(created_at__date=today).count()
    today_credit = (
        Transaction.objects.filter(created_at__date=today, transaction_type="credit").aggregate(total=Sum("amount"))[
            "total"
        ]
        or 0
    )
    today_debit = (
        Transaction.objects.filter(created_at__date=today, transaction_type="debit").aggregate(total=Sum("amount"))[
            "total"
        ]
        or 0
    )

    # Transaction trend - last 30 days (for bar chart)
    # Use 30-day window to ensure chart shows data even when recent activity is sparse
    trend_data = []
    for i in range(29, -1, -1):
        day = today - timedelta(days=i)
        day_credits = (
            Transaction.objects.filter(
                created_at__date=day, transaction_type__in=["credit", "interest", "dividend", "share_capital"]
            ).aggregate(total=Sum("amount"))["total"]
            or 0
        )
        day_debits = (
            Transaction.objects.filter(created_at__date=day, transaction_type__in=["debit", "transfer"]).aggregate(
                total=Sum("amount")
            )["total"]
            or 0
        )
        # Only include days that have transactions to keep chart clean
        if day_credits or day_debits:
            trend_data.append(
                {
                    "date": day.strftime("%b %d"),
                    "day": day.strftime("%a %d"),
                    "credits": float(day_credits),
                    "debits": float(day_debits),
                }
            )

    # If no data in 30 days, show last 7 days anyway (empty bars)
    if not trend_data:
        for i in range(6, -1, -1):
            day = today - timedelta(days=i)
            trend_data.append(
                {
                    "date": day.strftime("%b %d"),
                    "day": day.strftime("%a"),
                    "credits": 0,
                    "debits": 0,
                }
            )

    # Recent 10 transactions
    recent_transactions = (
        Transaction.objects.select_related("user", "member_account").all().order_by("-created_at")[:10]
    )

    # Overdue loan payments - get overdue repayments grouped by loan/member
    overdue_repayments = (
        LoanRepayment.objects.filter(payment_status="overdue")
        .select_related("loan_account", "loan_account__user", "loan_account__application")
        .order_by("due_date")
    )

    # Build overdue data: group by loan account, show earliest overdue + count
    overdue_loans = {}
    for rep in overdue_repayments:
        la_id = rep.loan_account_id
        la = rep.loan_account
        if la_id not in overdue_loans:
            overdue_loans[la_id] = {
                "loan_number": la.loan_number,
                "loan_id": la.id,
                "member_name": la.user.display_name,
                "member_id": la.user.member_id or "",
                "loan_type": la.get_loan_type_display(),
                "emi_amount": float(rep.amount_due),
                "oldest_due": rep.due_date.strftime("%b %d, %Y"),
                "days_overdue": (today - rep.due_date).days,
                "overdue_count": 0,
                "total_overdue": 0.0,
            }
        overdue_loans[la_id]["overdue_count"] += 1
        overdue_loans[la_id]["total_overdue"] += float(rep.amount_due)

    # Sort by days overdue (most overdue first), limit to 10
    overdue_list = sorted(overdue_loans.values(), key=lambda x: -x["days_overdue"])[:10]
    total_overdue_count = len(overdue_loans)

    # Fund stats
    fund_total_balance = FundAccount.objects.filter(is_deleted=False).aggregate(total=Sum("balance"))["total"] or 0
    total_funds = FundAccount.objects.filter(is_deleted=False).count()

    # Net profit for current fiscal year (April-March)
    from accounts.utils import get_financial_summary

    fw = effective_fy_window(request)
    financial_summary = get_financial_summary(fw.start_date, fw.end_date)
    net_profit = financial_summary["net_profit"]
    gross_revenue = financial_summary["gross_revenue"]

    return render(
        request,
        "admin/home.html",
        {
            "total_members": total_members,
            "active_members": active_members,
            "total_accounts": total_accounts,
            "active_accounts": active_accounts,
            "total_deposits": total_deposits,
            "total_transactions": total_transactions,
            "today_transactions": today_transactions,
            "total_receipts": total_transactions,
            "today_receipts": today_transactions,
            "today_credit": today_credit,
            "today_debit": today_debit,
            "recent_transactions": recent_transactions,
            "trend_data_json": json.dumps(trend_data),
            "account_type_data_json": json.dumps(account_type_data),
            "overdue_data_json": json.dumps(overdue_list),
            "total_overdue_count": total_overdue_count,
            "fund_total_balance": fund_total_balance,
            "total_funds": total_funds,
            "net_profit": net_profit,
            "gross_revenue": gross_revenue,
        },
    )


@login_required
@admin_required
def members_view(request):
    """Members management view with pagination and search"""
    page = request.GET.get("page", 1)
    query = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "").strip()
    member_type_filter = request.GET.get("member_type", "").strip()
    date_filter = request.GET.get("date_filter", "registered").strip() or "registered"
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    per_page = 25

    # Base queryset (exclude admin/staff accounts from member list)
    users = User.objects.filter(is_deleted=False, is_superuser=False, role="member").order_by("-date_joined")

    # Apply search filter if query provided
    if query:
        users = users.filter(
            Q(username__icontains=query)
            | Q(email__icontains=query)
            | Q(first_name__icontains=query)
            | Q(last_name__icontains=query)
            | Q(member_id__icontains=query)
            | Q(mobile_primary__icontains=query)
        )
    if status_filter:
        users = users.filter(status=status_filter)
    if member_type_filter:
        users = users.filter(member_type=member_type_filter)
    if date_filter not in ("registered", "resigned"):
        date_filter = "registered"
    if date_filter == "resigned":
        if date_from:
            users = users.filter(exit_date__gte=date_from)
        if date_to:
            users = users.filter(exit_date__lte=date_to)
    else:
        if date_from:
            users = users.filter(date_joined__date__gte=date_from)
        if date_to:
            users = users.filter(date_joined__date__lte=date_to)

    # Paginate
    paginator = Paginator(users, per_page)
    try:
        users_page = paginator.page(page)
    except PageNotAnInteger:
        users_page = paginator.page(1)
    except EmptyPage:
        users_page = paginator.page(paginator.num_pages)

    # Summary stats (exclude admin/staff accounts)
    total_members = User.objects.filter(is_deleted=False, is_superuser=False, role="member").count()
    active_members = User.objects.filter(is_deleted=False, is_superuser=False, role="member", status="active").count()

    member_status_choices = [
        ("active", "Active"),
        ("inactive", "Inactive"),
        ("resign", "Resigned"),
    ]

    return render(
        request,
        "admin/members.html",
        {
            "users": users_page,
            "search_query": query,
            "page_obj": users_page,
            "total_members": total_members,
            "active_members": active_members,
            "selected_status": status_filter,
            "selected_member_type": member_type_filter,
            "status_choices": member_status_choices,
            "member_type_choices": User.MEMBER_TYPE_CHOICES,
            "date_filter": date_filter,
            "date_from": date_from,
            "date_to": date_to,
        },
    )


@login_required
@admin_required
def add_member_view(request):
    """Create a new member with auto-generated password from DOB"""
    if request.method == "POST":
        try:
            # Rate limiting: max 10 member creations per minute per admin
            user_id = request.user.id
            now = timezone.now()
            if user_id in _member_creation_timestamps:
                recent_creates = [ts for ts in _member_creation_timestamps[user_id] if (now - ts).seconds < 60]
                if len(recent_creates) >= 10:
                    return JsonResponse({"success": False, "error": "Rate limit exceeded. Please wait a moment."})
                _member_creation_timestamps[user_id] = recent_creates
            else:
                _member_creation_timestamps[user_id] = []
            _member_creation_timestamps[user_id].append(now)

            username = request.POST.get("username", "").strip()
            email = request.POST.get("email", "").strip()
            full_name = request.POST.get("full_name", "").strip()
            date_of_birth = request.POST.get("date_of_birth", "").strip()
            mobile_primary = request.POST.get("mobile_primary", "").strip()
            role = request.POST.get("role", "member")

            if not username:
                return JsonResponse({"success": False, "error": "Username is required"})

            if not full_name:
                return JsonResponse({"success": False, "error": "Full name is required"})

            if not date_of_birth:
                return JsonResponse({"success": False, "error": "Date of birth is required"})

            # Parse date of birth
            try:
                dob = date.fromisoformat(date_of_birth)
            except ValueError:
                return JsonResponse({"success": False, "error": "Invalid date of birth format"})

            # Auto-generate password: DDMM + first 4 chars of name (uppercase, no spaces)
            name_chars = full_name.replace(" ", "").upper()[:4]
            password = f"{dob.day:02d}{dob.month:02d}{name_chars}"

            # Check for duplicate username
            if User.objects.filter(username=username).exists():
                return JsonResponse({"success": False, "error": "Username already exists"})

            # Create user
            user = User.objects.create_user(
                username=username,
                email=email,
                password=password,
                role=role,
            )

            # Set name fields
            first_name, last_name = split_full_name(full_name)
            user.first_name = first_name
            user.last_name = last_name
            user.date_of_birth = dob

            if mobile_primary:
                user.mobile_primary = mobile_primary

            if role == "admin":
                user.is_staff = True

            # Wrap member and account creation in transaction
            with transaction.atomic():
                # Auto-generate member_id: MBR-YEAR-SEQUENCE
                year = date.today().year
                last_member = (
                    User.objects.select_for_update()
                    .filter(member_id__startswith=f"MBR-{year}-")
                    .order_by("-member_id")
                    .first()
                )
                if last_member:
                    try:
                        last_seq = int(last_member.member_id.split("-")[-1])
                        next_seq = last_seq + 1
                    except (ValueError, IndexError):
                        next_seq = 1
                else:
                    next_seq = 1

                user.member_id = f"MBR-{year}-{next_seq:05d}"
                user.save()
                ensure_member_submodels(user)

                # Auto-create configured account types for the new member.
                account_label_map = dict(MemberAccount.ACCOUNT_TYPE_CHOICES)
                configured_rows = list(
                    AccountTypeConfiguration.objects.filter(is_active=True).order_by("display_order", "account_type")
                )
                account_types = []
                if configured_rows:
                    for cfg in configured_rows:
                        account_types.append(
                            (
                                cfg.account_type,
                                account_label_map.get(cfg.account_type, cfg.account_type.upper()),
                                cfg.interest_rate,
                            )
                        )
                else:
                    for account_type, account_label in MemberAccount.ACCOUNT_TYPE_CHOICES:
                        account_types.append(
                            (
                                account_type,
                                account_label,
                                _default_account_rates().get(account_type, Decimal("0.00")),
                            )
                        )

                created_accounts = []
                for account_type, account_name, default_rate in account_types:
                    # Generate account number: PREFIX-USERID-001
                    # Note: Using user.id ensures uniqueness since each user gets one account per configured type.
                    prefix_raw = "".join(ch for ch in str(account_type).upper() if ch.isalnum())
                    prefix = prefix_raw[:6] or "ACC"
                    account_number = f"{prefix}-{user.id:05d}-001"
                    account = MemberAccount.objects.create(
                        user=user,
                        account_number=account_number,
                        account_type=account_type,
                        status="active",
                        balance=0.00,
                        interest_rate=default_rate,
                        opening_date=date.today(),
                    )
                    created_accounts.append((account, account_name))

                # Log member creation
                log_action(
                    request,
                    "create",
                    "member",
                    user.id,
                    f"Created member {username} ({user.display_name}) with {len(created_accounts)} auto-generated accounts",
                )

                # Log each account creation
                for account, account_name in created_accounts:
                    log_action(
                        request,
                        "create",
                        "account",
                        account.id,
                        f"Auto-created {account_name} account {account.account_number} for member {username}",
                    )

                # Apply fund allocations for member registration (wrapped in try/except to never break member creation)
                try:
                    apply_fund_allocations(
                        request,
                        "member_registration",
                        Decimal("200.00"),
                        f"Member registration fee: {user.display_name} ({username})",
                        source_member=user,
                    )
                except Exception:
                    pass

            return JsonResponse({"success": True, "user_id": user.id, "password": password})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def delete_member_view(request, user_id):
    """Soft-delete a member (sets is_deleted flag instead of destroying data)"""
    if request.method == "POST":
        try:
            user = User.objects.get(id=user_id, is_deleted=False)
            # Prevent deleting yourself
            if user.id == request.user.id:
                return JsonResponse({"success": False, "error": "You cannot delete yourself"})

            user.is_deleted = True
            user.deleted_at = timezone.now()
            user.status = "closed"
            user.is_active = False
            user.save()

            # Also soft-delete all their accounts
            MemberAccount.objects.filter(user=user, is_deleted=False).update(
                is_deleted=True, deleted_at=timezone.now(), status="closed"
            )

            log_action(
                request,
                "delete",
                "member",
                user.id,
                f"Deleted member {user.username} ({user.display_name})",
            )

            return JsonResponse({"success": True})
        except User.DoesNotExist:
            return JsonResponse({"success": False, "error": "Member not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def add_account_view(request):
    """Add a new member account"""
    if request.method == "POST":
        try:
            user_id = request.POST.get("user_id")
            account_type = request.POST.get("account_type")
            balance = request.POST.get("balance", "0")
            interest_rate = request.POST.get("interest_rate", "").strip()
            principal_amount = request.POST.get("principal_amount", "0")
            opening_date = request.POST.get("opening_date")
            maturity_date = request.POST.get("maturity_date") or None
            tenure_months = request.POST.get("tenure_months") or None
            nominee_name = request.POST.get("nominee_name", "").strip()
            nominee_relationship = request.POST.get("nominee_relationship", "").strip()
            remarks = request.POST.get("remarks", "").strip()

            # Validate required fields
            if not user_id:
                return JsonResponse({"success": False, "error": "Member is required"})
            if not account_type:
                return JsonResponse({"success": False, "error": "Account type is required"})
            if not opening_date:
                return JsonResponse({"success": False, "error": "Opening date is required"})

            active_account_types = set(
                AccountTypeConfiguration.objects.filter(is_active=True).values_list("account_type", flat=True)
            )
            configured_rate = (
                AccountTypeConfiguration.objects.filter(account_type=account_type, is_active=True)
                .values_list("interest_rate", flat=True)
                .first()
            )
            if active_account_types and account_type not in active_account_types:
                return JsonResponse({"success": False, "error": "This account type is disabled in Settings"})
            if not interest_rate and configured_rate is not None:
                interest_rate = str(configured_rate)
            if not interest_rate:
                interest_rate = "0"

            # Validate user exists
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return JsonResponse({"success": False, "error": "Member not found"})

            # Auto-generate account number: PREFIX-YEAR-SEQUENCE
            type_prefix_map = {
                "fd": "FD",
                "cd": "CD",
                "rd": "RD",
                "od": "OD",
                "share": "SHR",
                "sukanya": "SKY",
                "suputra": "SPT",
            }
            prefix = type_prefix_map.get(account_type, "ACC")
            year = date.today().year

            # Convert tenure_months to int if provided
            if tenure_months:
                tenure_months = int(tenure_months)

            # Retry loop for duplicate account numbers
            max_retries = 5
            for attempt in range(max_retries):
                try:
                    with transaction.atomic():
                        # Lock the latest account row to prevent duplicate numbers
                        last_account = (
                            MemberAccount.objects.select_for_update()
                            .filter(account_number__startswith=f"{prefix}-{year}-")
                            .order_by("-account_number")
                            .first()
                        )
                        if last_account:
                            try:
                                last_seq = int(last_account.account_number.split("-")[-1])
                                next_seq = last_seq + 1
                            except (ValueError, IndexError):
                                next_seq = 1
                        else:
                            next_seq = 1

                        account_number = f"{prefix}-{year}-{next_seq:05d}"

                        account = MemberAccount.objects.create(
                            user=user,
                            account_number=account_number,
                            account_type=account_type,
                            balance=balance,
                            interest_rate=interest_rate,
                            principal_amount=principal_amount,
                            opening_date=opening_date,
                            maturity_date=maturity_date,
                            tenure_months=tenure_months,
                            nominee_name=nominee_name or None,
                            nominee_relationship=nominee_relationship or None,
                            remarks=remarks or None,
                        )
                    # If we get here, the transaction succeeded
                    log_action(
                        request,
                        "create",
                        "account",
                        account.id,
                        f"Created account {account_number} ({account_type}) for {user.display_name}",
                    )
                    # Phase B4: auto-post membership fee (once per user per FY).
                    from accounts.services import fees as fee_service

                    fee_service.apply_membership_fee(
                        user,
                        member_account=account,
                        actor=request.user,
                        ip_address=request.META.get("REMOTE_ADDR"),
                        audit_via="panel",
                    )
                    return JsonResponse({"success": True, "account_id": account.id})
                except IntegrityError:
                    if attempt == max_retries - 1:
                        raise
                    continue

        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def edit_account_view(request, account_id):
    """Edit a member account"""
    if request.method == "POST":
        try:
            account = MemberAccount.objects.get(id=account_id)

            account.account_type = request.POST.get("account_type", account.account_type)
            account.status = request.POST.get("status", account.status)
            account.balance = request.POST.get("balance", account.balance)
            account.interest_rate = request.POST.get("interest_rate", account.interest_rate)
            account.principal_amount = request.POST.get("principal_amount", account.principal_amount)

            opening_date = request.POST.get("opening_date")
            if opening_date:
                account.opening_date = opening_date

            maturity_date = request.POST.get("maturity_date")
            account.maturity_date = maturity_date or None

            tenure_months = request.POST.get("tenure_months")
            account.tenure_months = int(tenure_months) if tenure_months else None

            account.nominee_name = request.POST.get("nominee_name", "").strip() or None
            account.nominee_relationship = request.POST.get("nominee_relationship", "").strip() or None
            account.remarks = request.POST.get("remarks", "").strip() or None

            account.save()
            log_action(
                request,
                "update",
                "account",
                account.id,
                f"Updated account {account.account_number}",
            )
            return JsonResponse({"success": True})
        except MemberAccount.DoesNotExist:
            return JsonResponse({"success": False, "error": "Account not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def delete_member_account_view(request, account_id):
    """Soft-delete a member account (sets is_deleted flag instead of destroying data)"""
    if request.method == "POST":
        try:
            account = MemberAccount.objects.get(id=account_id, is_deleted=False)
            account.is_deleted = True
            account.deleted_at = timezone.now()
            account.status = "closed"
            account.save()
            log_action(
                request,
                "delete",
                "account",
                account.id,
                f"Deleted account {account.account_number} ({account.get_account_type_display()}) "
                f"for {account.user.display_name}",
            )
            return JsonResponse({"success": True})
        except MemberAccount.DoesNotExist:
            return JsonResponse({"success": False, "error": "Account not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def search_members_view(request):
    """Search members for autocomplete"""
    query = request.GET.get("q", "").strip()
    if not query or len(query) < 2:
        return JsonResponse({"results": []})

    users = User.objects.filter(
        Q(username__icontains=query)
        | Q(first_name__icontains=query)
        | Q(last_name__icontains=query)
        | Q(member_id__icontains=query),
        is_deleted=False,
        is_superuser=False,
        role="member",
    )[:10]

    results = [
        {
            "id": user.id,
            "username": user.username,
            "full_name": user.display_name,
            "member_id": user.member_id or "",
        }
        for user in users
    ]

    return JsonResponse({"results": results})


@login_required
@admin_required
def get_member_view(request, user_id):
    """Get member data for viewing and editing"""
    try:
        user = User.objects.get(id=user_id)

        current_address = format_current_address_display(user)
        permanent_address = format_permanent_address_display(user)

        kyc = user_kyc(user)
        primary_nom = nominee_primary(user)
        alt_nom = nominee_secondary(user)
        share_lot = user_primary_share_lot(user)
        sc_total = user_active_share_capital_total(user)
        n_shares = user_active_share_count(user)

        # Get member accounts summary
        accounts = MemberAccount.objects.filter(user=user, is_deleted=False).order_by("account_type")
        accounts_data = [
            {
                "id": a.id,
                "account_number": a.account_number,
                "account_type": a.account_type,
                "account_type_display": a.get_account_type_display(),
                "status": a.status,
                "status_display": a.get_status_display(),
                "balance": str(a.balance),
                "interest_rate": str(a.interest_rate),
                "opening_date": a.opening_date.strftime("%b %d, %Y") if a.opening_date else "",
                "maturity_date": a.maturity_date.strftime("%b %d, %Y") if a.maturity_date else "",
            }
            for a in accounts
        ]

        return JsonResponse(
            {
                "success": True,
                "user": {
                    "id": user.id,
                    "username": user.username,
                    "email": user.email,
                    # Core Identity
                    "member_id": user.member_id or "",
                    "full_name": user.display_name,
                    "member_type": user.get_member_type_display() if user.member_type else "",
                    "member_type_raw": user.member_type or "",
                    "status": user.get_status_display() if user.status else "",
                    "status_raw": user.status or "",
                    "date_of_joining": user.date_of_joining.strftime("%b %d, %Y") if user.date_of_joining else "",
                    "inactive_since": user.inactive_since.strftime("%b %d, %Y") if user.inactive_since else "",
                    # Personal
                    "date_of_birth": user.date_of_birth.strftime("%b %d, %Y") if user.date_of_birth else "",
                    "age": user.age or "",
                    "gender": user.get_gender_display() if user.gender else "",
                    "marital_status": user.get_marital_status_display() if user.marital_status else "",
                    "occupation": user.occupation or "",
                    "annual_income_bracket": user.annual_income_bracket or "",
                    "education_qualification": user.education_qualification or "",
                    # Contact
                    "mobile_primary": user.mobile_primary or "",
                    "mobile_alternate": user.mobile_alternate or "",
                    "preferred_comm_mode": user.get_preferred_comm_mode_display() if user.preferred_comm_mode else "",
                    "dnd_enabled": user.dnd_enabled,
                    # Address
                    "current_address": current_address or "",
                    "permanent_address": permanent_address or "",
                    # KYC
                    "kyc_status": kyc.get_kyc_status_display() if kyc else "",
                    "kyc_verified_date": kyc.kyc_verified_date.strftime("%b %d, %Y")
                    if kyc and kyc.kyc_verified_date
                    else "",
                    "aadhar_number": kyc.aadhaar_number or "" if kyc else "",
                    "pan_number": kyc.pan_number or "" if kyc else "",
                    "voter_id": kyc.voter_id or "" if kyc else "",
                    "passport_number": kyc.passport_number or "" if kyc else "",
                    "driving_license": kyc.driving_licence or "" if kyc else "",
                    # Nominee
                    "nominee_name": primary_nom.name if primary_nom else "",
                    "nominee_relationship": primary_nom.relationship or "" if primary_nom else "",
                    "nominee_dob": primary_nom.dob.strftime("%b %d, %Y") if primary_nom and primary_nom.dob else "",
                    "nominee_contact": primary_nom.contact or "" if primary_nom else "",
                    "nominee_address": primary_nom.address or "" if primary_nom else "",
                    "alt_nominee_name": alt_nom.name if alt_nom else "",
                    "alt_nominee_relationship": alt_nom.relationship or "" if alt_nom else "",
                    # Shareholding
                    "share_capital_amount": str(sc_total),
                    "number_of_shares": n_shares,
                    "face_value_per_share": str(share_lot.face_value_per_share) if share_lot else "0.00",
                    "share_certificate_number": share_lot.certificate_number or "" if share_lot else "",
                    "share_issue_date": share_lot.issue_date.strftime("%b %d, %Y")
                    if share_lot and share_lot.issue_date
                    else "",
                    "dividend_payable_balance": str(user.dividend_payable_balance),
                    "last_dividend_paid_date": user.last_dividend_paid_date.strftime("%b %d, %Y")
                    if user.last_dividend_paid_date
                    else "",
                    # Banking flags
                    "eligible_for_accounts": user.eligible_for_accounts,
                    "eligible_for_loans": user.eligible_for_loans,
                    "eligible_for_dividend": user.eligible_for_dividend,
                    "eligible_for_voting": user.eligible_for_voting,
                    "risk_category": user.get_risk_category_display() if user.risk_category else "",
                    # Internal
                    "internal_remarks": user.internal_remarks or "",
                    "closure_reason": user.closure_reason or "",
                    "aml_check_status": user.aml_check_status or "",
                    # Accounts summary
                    "accounts": accounts_data,
                    "accounts_count": len(accounts_data),
                    # Edit modal fields
                    "first_name": user.first_name,
                    "last_name": user.last_name,
                    "role": user.role,
                    "role_display": user.get_role_display_name(),
                    "is_active": user.is_active,
                    "is_staff": user.is_staff,
                    "date_joined": user.date_joined.strftime("%b %d, %Y") if user.date_joined else "",
                    "last_login": user.last_login.strftime("%b %d, %Y %H:%M") if user.last_login else "Never",
                },
            }
        )
    except User.DoesNotExist:
        return JsonResponse({"success": False, "error": "User not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def edit_member_view(request, user_id):
    """Edit a member"""
    if request.method == "POST":
        try:
            user = User.objects.get(id=user_id)

            username = request.POST.get("username")
            email = request.POST.get("email")
            full_name = request.POST.get("full_name", "").strip() or request.POST.get("name", "").strip()
            mobile_primary = request.POST.get("mobile_primary", "")
            role = request.POST.get("role", "member")
            password = request.POST.get("password", "").strip()
            status = request.POST.get("status", user.status).strip() or user.status
            closure_reason = request.POST.get("closure_reason", "").strip()

            # Validate
            if not all([username, email]):
                return JsonResponse({"success": False, "error": "Username and email are required"})

            # Check if username is taken by another user
            if User.objects.filter(username=username).exclude(id=user_id).exists():
                return JsonResponse({"success": False, "error": "Username already exists"})

            # Check if email is taken by another user
            if User.objects.filter(email=email).exclude(id=user_id).exists():
                return JsonResponse({"success": False, "error": "Email already exists"})

            # Validate role
            valid_roles = ["admin", "member", "viewer"]
            if role not in valid_roles:
                role = "member"
            valid_statuses = ["active", "inactive", "resign"]
            if status not in valid_statuses:
                status = user.status

            # Split name into first_name and last_name using utility
            first_name, last_name = split_full_name(full_name)
            user.first_name = first_name
            user.last_name = last_name

            # Update user
            user.username = username
            user.email = email
            user.mobile_primary = mobile_primary
            user.role = role
            user.is_staff = role == "admin"
            previous_status = user.status
            user.status = status
            if closure_reason:
                user.closure_reason = closure_reason

            # Update password if provided (with validation)
            if password:
                is_valid, errors = validate_password_strength(password)
                if not is_valid:
                    return JsonResponse({"success": False, "error": errors[0]})
                user.set_password(password)

            with transaction.atomic():
                _apply_member_status_rules(user, status, previous_status=previous_status)
                user.save()

                settlement_transactions = []
                if status == "resign":
                    settlement_accounts = MemberAccount.objects.select_for_update().filter(
                        user=user,
                        is_deleted=False,
                        status="active",
                        balance__gt=0,
                    )
                    for account in settlement_accounts:
                        line = {
                            "account_id": account.id,
                            "transaction_type": "debit",
                            "amount": str(account.balance),
                            "description": "Member resignation settlement",
                        }
                        settlement_transactions.append(
                            _process_transaction_line(
                                user=user,
                                line=line,
                                payment_mode="cash",
                                reference_number="",
                                remarks="Auto settlement on resignation",
                                created_by=request.user,
                            )
                        )

            log_action(
                request,
                "update",
                "member",
                user.id,
                f"Updated member {user.username} ({user.display_name}) with status {user.status}",
            )
            if status == "resign" and settlement_transactions:
                log_action(
                    request,
                    "create",
                    "transaction",
                    settlement_transactions[-1].id,
                    f"Generated {len(settlement_transactions)} resignation settlement transactions for {user.display_name}",
                )
            return JsonResponse(
                {
                    "success": True,
                    "settlement_transactions": [r.transaction_number for r in settlement_transactions],
                }
            )
        except User.DoesNotExist:
            return JsonResponse({"success": False, "error": "User not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def reset_member_password_view(request, user_id):
    """Reset member password to DOB-based format"""
    if request.method == "POST":
        try:
            user = User.objects.get(id=user_id)

            if not user.date_of_birth or not user.first_name:
                return JsonResponse(
                    {"success": False, "error": "Cannot reset password: Date of birth or name is missing"}
                )

            new_password = user.reset_password_to_dob()
            if new_password:
                user.save()
                log_action(
                    request,
                    "reset_password",
                    "member",
                    user.id,
                    f"Reset password for {user.username} ({user.display_name})",
                )
                return JsonResponse(
                    {"success": True, "password": new_password, "message": f"Password reset to: {new_password}"}
                )
            else:
                return JsonResponse({"success": False, "error": "Failed to generate password"})

        except User.DoesNotExist:
            return JsonResponse({"success": False, "error": "User not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def profile_view(request):
    """Admin profile view"""
    if request.method == "POST":
        # Handle name field - split into first_name and last_name using utility
        name = request.POST.get("name", "").strip()
        first_name, last_name = split_full_name(name)
        request.user.first_name = first_name
        request.user.last_name = last_name

        request.user.email = request.POST.get("email", "")
        request.user.mobile_primary = request.POST.get("mobile_primary", "")
        request.user.save()
        messages.success(request, "Profile updated successfully!")
        return redirect("/profile/")

    return render(request, "admin/profile.html")


@login_required
@admin_required
def accounts_view(request):
    """Account Book view - shows all members and their account balances in a grid"""
    query = request.GET.get("q", "").strip()
    account_type_filter = request.GET.get("account_type", "").strip()
    account_status_filter = request.GET.get("account_status", "").strip()
    member_status_filter = request.GET.get("member_status", "").strip()
    date_filter = request.GET.get("date_filter", "member_registered").strip() or "member_registered"
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    page = request.GET.get("page", 1)
    per_page = 50

    # Get all members with their accounts
    members = User.objects.filter(role="member", is_deleted=False).order_by("first_name", "last_name")

    # Apply search filter
    if query:
        members = members.filter(
            Q(first_name__icontains=query)
            | Q(last_name__icontains=query)
            | Q(username__icontains=query)
            | Q(member_id__icontains=query)
        )
    if member_status_filter:
        members = members.filter(status=member_status_filter)
    if date_filter not in ("member_registered", "account_opened"):
        date_filter = "member_registered"
    if date_filter == "member_registered":
        if date_from:
            members = members.filter(date_joined__date__gte=date_from)
        if date_to:
            members = members.filter(date_joined__date__lte=date_to)

    # Paginate members
    paginator = Paginator(members, per_page)
    try:
        members_page = paginator.page(page)
    except PageNotAnInteger:
        members_page = paginator.page(1)
    except EmptyPage:
        members_page = paginator.page(paginator.num_pages)

    # Build account book data structure
    # Only show account types that are configured as active in Settings > Products.
    # Use the custom product name from the database as the column header.
    from accounts.models import AccountTypeConfiguration

    active_configs = AccountTypeConfiguration.objects.filter(is_active=True).order_by("display_order", "account_type")
    configured_types = set(c.account_type.lower() for c in active_configs)

    # Build account_types using the operator's custom names from the DB.
    # Map each configured product to the matching ACCOUNT_TYPE_CHOICES code.
    all_choices = MemberAccount.ACCOUNT_TYPE_CHOICES
    choices_code_map = {code.lower(): code for code, _ in all_choices}
    account_types = []
    for cfg in active_configs:
        code = choices_code_map.get(cfg.account_type.lower())
        if code:
            account_types.append((code, cfg.account_type))  # Use the DB name as label

    # Fallback: if nothing is configured yet (fresh install), show all types.
    if not account_types:
        account_types = list(all_choices)

    account_book = []

    for member in members_page:
        member_data = {
            "id": member.id,
            "name": member.display_name,
            "member_id": member.member_id or "-",
            "username": member.username,
            "account_balances": [],  # List of balances in order
            "total_balance": 0,
        }

        # Get all accounts for this member (only configured types)
        accounts = MemberAccount.objects.filter(user=member, is_deleted=False)
        if configured_types:
            accounts = accounts.filter(account_type__in=[c.upper() for c in configured_types] + list(configured_types))
        if account_type_filter:
            accounts = accounts.filter(account_type=account_type_filter)
        if account_status_filter:
            accounts = accounts.filter(status=account_status_filter)
        if date_filter == "account_opened":
            if date_from:
                accounts = accounts.filter(opening_date__gte=date_from)
            if date_to:
                accounts = accounts.filter(opening_date__lte=date_to)
        accounts_dict = {acc.account_type: acc for acc in accounts}

        # Build ordered list of balances matching account_types order
        for account_type, _ in account_types:
            if account_type in accounts_dict:
                acc = accounts_dict[account_type]
                member_data["account_balances"].append(
                    {
                        "balance": acc.balance,
                        "status": acc.status,
                    }
                )
                if acc.status == "active":
                    member_data["total_balance"] += acc.balance
            else:
                member_data["account_balances"].append(
                    {
                        "balance": 0,
                        "status": "inactive",
                    }
                )

        account_book.append(member_data)

    # Calculate totals for each account type
    account_totals = []
    grand_total = 0
    totals_qs = MemberAccount.objects.filter(is_deleted=False)
    if configured_types:
        totals_qs = totals_qs.filter(account_type__in=[c.upper() for c in configured_types] + list(configured_types))
    if account_status_filter:
        totals_qs = totals_qs.filter(status=account_status_filter)
    if member_status_filter:
        totals_qs = totals_qs.filter(user__status=member_status_filter)
    if date_filter == "member_registered":
        if date_from:
            totals_qs = totals_qs.filter(user__date_joined__date__gte=date_from)
        if date_to:
            totals_qs = totals_qs.filter(user__date_joined__date__lte=date_to)
    else:
        if date_from:
            totals_qs = totals_qs.filter(opening_date__gte=date_from)
        if date_to:
            totals_qs = totals_qs.filter(opening_date__lte=date_to)
    for account_type, _ in account_types:
        type_qs = totals_qs.filter(account_type=account_type)
        total = type_qs.aggregate(total=Sum("balance"))["total"] or 0
        account_totals.append(total)
        grand_total += total

    return render(
        request,
        "admin/accounts.html",
        {
            "account_book": account_book,
            "account_types": account_types,
            "account_totals": account_totals,
            "grand_total": grand_total,
            "search_query": query,
            "selected_account_type": account_type_filter,
            "selected_account_status": account_status_filter,
            "selected_member_status": member_status_filter,
            "account_status_choices": MemberAccount.STATUS_CHOICES,
            "member_status_choices": [("active", "Active"), ("inactive", "Inactive"), ("resign", "Resigned")],
            "page_obj": members_page,
            "date_filter": date_filter,
            "date_from": date_from,
            "date_to": date_to,
        },
    )


@login_required
@admin_required
def get_account_view(request, account_id):
    """Get account data for viewing/editing"""
    try:
        account = MemberAccount.objects.select_related("user").get(id=account_id)
        return JsonResponse(
            {
                "success": True,
                "account": {
                    "id": account.id,
                    "account_number": account.account_number,
                    "account_type": account.account_type,
                    "account_type_display": account.get_account_type_display(),
                    "status": account.status,
                    "status_display": account.get_status_display(),
                    "balance": str(account.balance),
                    "interest_rate": str(account.interest_rate),
                    "principal_amount": str(account.principal_amount),
                    "opening_date": account.opening_date.strftime("%Y-%m-%d") if account.opening_date else "",
                    "maturity_date": account.maturity_date.strftime("%Y-%m-%d") if account.maturity_date else "",
                    "tenure_months": account.tenure_months or "",
                    "user_id": account.user.id,
                    "user_name": account.user.display_name,
                    "member_id": account.user.member_id or "",
                    "nominee_name": account.nominee_name or "",
                    "nominee_relationship": account.nominee_relationship or "",
                    "remarks": account.remarks or "",
                },
            }
        )
    except MemberAccount.DoesNotExist:
        return JsonResponse({"success": False, "error": "Account not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def get_account_transactions_view(request, account_id):
    """Get paginated transaction history for an account"""
    try:
        account = MemberAccount.objects.get(id=account_id)
        page = int(request.GET.get("page", 1))
        per_page = 20

        transactions = Transaction.objects.filter(member_account=account).order_by("-created_at")
        total = transactions.count()
        start = (page - 1) * per_page
        end = start + per_page
        page_transactions = transactions[start:end]

        txns = []
        for r in page_transactions:
            txns.append(
                {
                    "transaction_number": r.transaction_number,
                    "date": r.created_at.strftime("%b %d, %Y") if r.created_at else "-",
                    "type": r.get_transaction_type_display(),
                    "type_raw": r.transaction_type,
                    "description": r.description or "-",
                    "amount": str(r.amount),
                    "balance_after": str(r.balance_after) if r.balance_after is not None else "-",
                }
            )

        return JsonResponse(
            {
                "success": True,
                "transactions": txns,
                "has_more": end < total,
                "total": total,
                "page": page,
            }
        )
    except MemberAccount.DoesNotExist:
        return JsonResponse({"success": False, "error": "Account not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def account_details_view(request, account_id):
    """HTML account details page to avoid exposing raw JSON route."""
    try:
        account = MemberAccount.objects.select_related("user").get(id=account_id, is_deleted=False)
    except MemberAccount.DoesNotExist:
        messages.error(request, "Account not found.")
        return redirect("/accounts/")
    transactions = Transaction.objects.filter(member_account=account).order_by("-created_at")[:50]
    return render(
        request,
        "admin/account_details.html",
        {
            "account": account,
            "transactions": transactions,
        },
    )


# ========================================
# Transaction Views
# ========================================


@login_required
@admin_required
def transactions_view(request):
    """Transaction list view with pagination, search, and filters"""
    page = request.GET.get("page", 1)
    query = request.GET.get("q", "").strip()
    transaction_type = request.GET.get("type", "").strip()
    payment_mode = request.GET.get("mode", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    per_page = 25

    # Base queryset with related data
    receipts = Transaction.objects.select_related("user", "member_account", "created_by").all().order_by("-created_at")

    # Apply search filter
    if query:
        receipts = receipts.filter(
            Q(transaction_number__icontains=query)
            | Q(user__first_name__icontains=query)
            | Q(user__last_name__icontains=query)
            | Q(user__member_id__icontains=query)
            | Q(member_account__account_number__icontains=query)
            | Q(reference_number__icontains=query)
        )

    # Apply type filter
    if transaction_type:
        receipts = receipts.filter(transaction_type=transaction_type)

    # Apply payment mode filter
    if payment_mode:
        receipts = receipts.filter(payment_mode=payment_mode)
    if date_from:
        receipts = receipts.filter(created_at__date__gte=date_from)
    if date_to:
        receipts = receipts.filter(created_at__date__lte=date_to)

    # Paginate
    paginator = Paginator(receipts, per_page)
    try:
        receipts_page = paginator.page(page)
    except PageNotAnInteger:
        receipts_page = paginator.page(1)
    except EmptyPage:
        receipts_page = paginator.page(paginator.num_pages)

    # Summary stats
    total_receipts = Transaction.objects.count()
    today_receipts = Transaction.objects.filter(created_at__date=timezone.now().date()).count()

    return render(
        request,
        "admin/transactions.html",
        {
            "receipts": receipts_page,
            "search_query": query,
            "selected_type": transaction_type,
            "selected_mode": payment_mode,
            "date_from": date_from,
            "date_to": date_to,
            "transaction_type_choices": Transaction.TRANSACTION_TYPE_CHOICES,
            "payment_mode_choices": Transaction.PAYMENT_MODE_CHOICES,
            "page_obj": receipts_page,
            "total_receipts": total_receipts,
            "today_receipts": today_receipts,
            "society_name": getattr(settings, "SOCIETY_NAME", "Panels"),
            "funds": FundAccount.objects.filter(is_deleted=False, is_active=True).order_by("name"),
            "fund_options_json": json.dumps(
                [
                    {"id": item.id, "name": item.name, "type": item.get_fund_type_display()}
                    for item in FundAccount.objects.filter(is_deleted=False, is_active=True).order_by("name")
                ]
            ),
        },
    )


def _instrument_json_for_receipt_modal(inst):
    """Serialize a linked Instrument for the receipt modal JSON (Phase A2)."""
    if not inst:
        return None
    return {
        "id": inst.id,
        "instrument_type": inst.instrument_type,
        "instrument_type_display": inst.get_instrument_type_display(),
        "amount": str(inst.amount),
        "cheque_number": inst.cheque_number or "",
        "drawer_name": inst.drawer_name or "",
        "drawer_bank": inst.drawer_bank or "",
        "drawer_ifsc": inst.drawer_ifsc or "",
        "cheque_date": inst.cheque_date.isoformat() if inst.cheque_date else "",
        "cheque_status": inst.cheque_status or "",
        "reference_number": inst.reference_number or "",
        "upi_vpa": inst.upi_vpa or "",
        "is_cleared": inst.is_cleared,
        "clearing_date": inst.clearing_date.isoformat() if inst.clearing_date else "",
        "bounce_reason": inst.bounce_reason or "",
    }


@login_required
@admin_required
def get_transaction_view(request, transaction_id):
    """Get transaction data for viewing."""
    try:
        receipt = Transaction.objects.select_related("user", "member_account", "created_by", "instrument").get(
            id=transaction_id
        )
        return JsonResponse(
            {
                "success": True,
                "receipt": {
                    "id": receipt.id,
                    "receipt_number": receipt.transaction_number,
                    "transaction_type": receipt.transaction_type,
                    "transaction_type_display": receipt.get_transaction_type_display(),
                    "amount": str(receipt.amount),
                    "description": receipt.description or "",
                    "payment_mode": receipt.payment_mode,
                    "payment_mode_display": receipt.get_payment_mode_display(),
                    "reference_number": receipt.reference_number or "",
                    "balance_after": str(receipt.balance_after),
                    "remarks": receipt.remarks or "",
                    "created_at": receipt.created_at.strftime("%b %d, %Y %I:%M %p") if receipt.created_at else "",
                    "created_date": receipt.created_at.strftime("%b %d, %Y") if receipt.created_at else "",
                    "instrument": _instrument_json_for_receipt_modal(receipt.instrument),
                    # Member info
                    "member_name": receipt.user.display_name,
                    "member_id": receipt.user.member_id or "",
                    "member_mobile": receipt.user.mobile_primary or "",
                    "member_address": format_current_address_display(receipt.user),
                    # Account info
                    "account_number": receipt.member_account.account_number,
                    "account_type": receipt.member_account.account_type,
                    "account_type_display": receipt.member_account.get_account_type_display(),
                    # Admin info
                    "created_by_name": receipt.created_by.display_name if receipt.created_by else "",
                },
            }
        )
    except Transaction.DoesNotExist:
        return JsonResponse({"success": False, "error": "Transaction not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def add_transaction_view(request):
    """Create one or more transactions, or stage voucher entries (delegates to
    `services.transactions.post_voucher` / `post_transactions_bulk`)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})

    user_id = request.POST.get("user_id")
    if not user_id:
        return JsonResponse({"success": False, "error": "Member is required"})

    payment_mode = request.POST.get("payment_mode", "cash")
    reference_number = request.POST.get("reference_number", "").strip()
    remarks = request.POST.get("remarks", "").strip()
    use_voucher = request.POST.get("use_voucher", "").lower() in ("1", "true", "yes", "on")
    # A4: vouchers are opt-in; the operator chooses via the use_voucher checkbox.

    voucher_type = request.POST.get("voucher_type", "receipt").strip() or "receipt"
    fund_id = request.POST.get("fund_id", "").strip() or None
    account_entries_raw = request.POST.get("account_entries", "").strip()
    if not account_entries_raw:
        return JsonResponse({"success": False, "error": "At least one account entry is required"})

    try:
        account_entries = json.loads(account_entries_raw)
    except json.JSONDecodeError:
        return JsonResponse({"success": False, "error": "Invalid account entries payload"})

    instrument_payload = None
    raw_inst = request.POST.get("instrument_payload", "").strip()
    if raw_inst:
        try:
            parsed_inst = json.loads(raw_inst)
        except json.JSONDecodeError:
            return JsonResponse({"success": False, "error": "Invalid instrument_payload JSON"})
        if not isinstance(parsed_inst, dict):
            return JsonResponse({"success": False, "error": "instrument_payload must be a JSON object"})
        instrument_payload = parsed_inst

    ip = _get_client_ip(request)

    try:
        if use_voucher:
            voucher = transaction_service.post_voucher(
                user_id=user_id,
                voucher_type=voucher_type,
                lines=account_entries,
                payment_mode=payment_mode,
                reference_number=reference_number,
                remarks=remarks,
                instrument_payload=instrument_payload,
                actor=request.user,
                ip_address=ip,
                audit_via="panel",
            )
            return JsonResponse({"success": True, "voucher_id": voucher.id, "voucher_number": voucher.voucher_number})

        result = transaction_service.post_transactions_bulk(
            user_id=user_id,
            lines=account_entries,
            payment_mode=payment_mode,
            reference_number=reference_number,
            remarks=remarks,
            fund_id=fund_id,
            instrument_payload=instrument_payload,
            actor=request.user,
            ip_address=ip,
            audit_via="panel",
        )
    except ServiceError as e:
        return JsonResponse({"success": False, "error": e.message})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})

    created_receipts = result["transactions"]

    # Best-effort loan-interest fund allocation per interest-type line.
    # Adapter concern for now because `apply_fund_allocations` still takes a
    # `request` (will move into the service layer when Phase B4 overhauls fees).
    for receipt in created_receipts:
        if receipt.transaction_type == "interest":
            try:
                apply_fund_allocations(
                    request,
                    "loan_interest",
                    receipt.amount,
                    f"Loan interest: {receipt.user.display_name} - {receipt.member_account.account_number}",
                    source_member=receipt.user,
                )
            except Exception:
                pass

    return JsonResponse(
        {
            "success": True,
            "receipt_id": created_receipts[0].id,
            "receipt_ids": [item.id for item in created_receipts],
        }
    )


@login_required
@admin_required
def vouchers_view(request):
    """List pending and transferred vouchers for receipts page."""
    vouchers = Voucher.objects.select_related("user", "created_by", "transferred_to_fund", "instrument").prefetch_related(
        "entries"
    )[:100]
    data = []
    for voucher in vouchers:
        data.append(
            {
                "id": voucher.id,
                "voucher_number": voucher.voucher_number,
                "member_name": voucher.user.display_name,
                "member_id": voucher.user.member_id or "",
                "total_amount": str(voucher.total_amount),
                "status": voucher.status,
                "status_display": voucher.get_status_display(),
                "voucher_type": voucher.voucher_type,
                "voucher_type_display": voucher.get_voucher_type_display(),
                "line_count": voucher.entries.count(),
                "created_at": voucher.created_at.strftime("%d %b %Y, %H:%M"),
                "fund_name": voucher.transferred_to_fund.name if voucher.transferred_to_fund else "",
                "payment_mode": voucher.payment_mode,
                "payment_mode_display": voucher.get_payment_mode_display(),
                "instrument_type": voucher.instrument.get_instrument_type_display() if voucher.instrument else "",
            }
        )
    return JsonResponse({"success": True, "vouchers": data})


@login_required
@admin_required
def vouchers_page_view(request):
    """Dedicated vouchers listing page with search, filters, and pagination."""
    page = request.GET.get("page", 1)
    query = request.GET.get("q", "").strip()
    voucher_type = request.GET.get("voucher_type", "").strip()
    status = request.GET.get("status", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    per_page = 25

    vouchers_qs = Voucher.objects.select_related(
        "user", "created_by", "transferred_to_fund", "instrument"
    ).prefetch_related("entries").order_by("-created_at")

    if query:
        vouchers_qs = vouchers_qs.filter(
            Q(voucher_number__icontains=query)
            | Q(user__first_name__icontains=query)
            | Q(user__last_name__icontains=query)
            | Q(user__member_id__icontains=query)
        )

    if voucher_type:
        vouchers_qs = vouchers_qs.filter(voucher_type=voucher_type)

    if status:
        vouchers_qs = vouchers_qs.filter(status=status)

    if date_from:
        vouchers_qs = vouchers_qs.filter(created_at__date__gte=date_from)

    if date_to:
        vouchers_qs = vouchers_qs.filter(created_at__date__lte=date_to)

    paginator = Paginator(vouchers_qs, per_page)
    try:
        vouchers_page = paginator.page(page)
    except PageNotAnInteger:
        vouchers_page = paginator.page(1)
    except EmptyPage:
        vouchers_page = paginator.page(paginator.num_pages)

    total_vouchers = Voucher.objects.count()
    pending_vouchers = Voucher.objects.filter(status="pending").count()

    return render(
        request,
        "admin/vouchers.html",
        {
            "vouchers": vouchers_page,
            "search_query": query,
            "selected_voucher_type": voucher_type,
            "selected_status": status,
            "date_from": date_from,
            "date_to": date_to,
            "page_obj": vouchers_page,
            "total_vouchers": total_vouchers,
            "pending_vouchers": pending_vouchers,
            "funds": FundAccount.objects.filter(is_deleted=False, is_active=True).order_by("name"),
            "active": "vouchers",
        },
    )


@login_required
@admin_required
def transfer_voucher_to_fund_view(request, voucher_id):
    """Convert voucher lines to receipts and move total to a fund."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})
    fund_id = request.POST.get("fund_id", "").strip()
    if not fund_id:
        return JsonResponse({"success": False, "error": "Fund is required"})
    try:
        with transaction.atomic():
            voucher = Voucher.objects.select_for_update().select_related("instrument").get(id=voucher_id)
            if voucher.status != "pending":
                return JsonResponse({"success": False, "error": "Voucher is already processed"})
            fund = FundAccount.objects.select_for_update().get(id=fund_id, is_deleted=False, is_active=True)
            receipts = []
            shared_inst = voucher.instrument
            for entry in voucher.entries.select_related("member_account").all():
                line = {
                    "account_id": entry.member_account_id,
                    "transaction_type": entry.transaction_type,
                    "amount": str(entry.amount),
                    "description": entry.description,
                    "loan_repayment_id": entry.linked_loan_repayment_id,
                }
                receipt = _process_transaction_line(
                    user=voucher.user,
                    line=line,
                    payment_mode=voucher.payment_mode,
                    reference_number=voucher.reference_number or "",
                    remarks=voucher.remarks or "",
                    created_by=request.user,
                    instrument=shared_inst,
                )
                entry.created_transaction = receipt
                entry.save(update_fields=["created_transaction"])
                receipts.append(receipt)
            FundAccount.objects.filter(id=fund.id).update(balance=F("balance") + voucher.total_amount)
            fund.refresh_from_db(fields=["balance"])
            FundTransaction.objects.create(
                fund=fund,
                transaction_type="credit",
                amount=voucher.total_amount,
                description=f"Voucher transfer {voucher.voucher_number}",
                payment_mode="internal",
                reference_number=voucher.reference_number or None,
                balance_after=fund.balance,
                source_member=voucher.user,
                created_by=request.user,
            )
            voucher.status = "transferred"
            voucher.transferred_to_fund = fund
            voucher.transferred_at = timezone.now()
            voucher.save(update_fields=["status", "transferred_to_fund", "transferred_at"])
        log_action(
            request,
            "create",
            "voucher",
            voucher.id,
            f"Transferred {voucher.get_voucher_type_display()} {voucher.voucher_number} to fund {fund.name}",
        )
        return JsonResponse({"success": True, "receipt_ids": [item.id for item in receipts]})
    except Voucher.DoesNotExist:
        return JsonResponse({"success": False, "error": "Voucher not found"})
    except FundAccount.DoesNotExist:
        return JsonResponse({"success": False, "error": "Fund not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def get_member_accounts_view(request, user_id):
    """Get accounts for a specific member (for cascading dropdown)"""
    try:
        user = User.objects.get(id=user_id)
        accounts = MemberAccount.objects.filter(user=user, status="active", is_deleted=False).order_by("account_type")

        results = [
            {
                "id": account.id,
                "account_number": account.account_number,
                "account_type": account.account_type,
                "account_type_display": account.get_account_type_display(),
                "balance": str(account.balance),
            }
            for account in accounts
        ]

        pending_repayments = (
            LoanRepayment.objects.filter(loan_account__user=user, loan_account__status="active")
            .exclude(payment_status="paid")
            .select_related("loan_account", "loan_account__disbursement_account")
            .order_by("due_date")[:20]
        )
        pending_emi_data = [
            {
                "id": repayment.id,
                "loan_id": repayment.loan_account.id,
                "loan_number": repayment.loan_account.loan_number,
                "installment_number": repayment.installment_number,
                "amount_due": str(repayment.amount_due),
                "due_date": repayment.due_date.strftime("%Y-%m-%d"),
                "status": repayment.payment_status,
                "account_id": repayment.loan_account.disbursement_account_id or "",
            }
            for repayment in pending_repayments
        ]
        return JsonResponse({"success": True, "accounts": results, "pending_emis": pending_emi_data})
    except User.DoesNotExist:
        return JsonResponse({"success": False, "error": "Member not found"})


@login_required
@admin_required
def get_member_transactions_view(request, user_id):
    """Get recent transactions for a member (for modal transaction history tab)"""
    try:
        user = User.objects.get(id=user_id)
        page = int(request.GET.get("page", 1))
        per_page = 15

        receipts = Transaction.objects.select_related("member_account").filter(user=user).order_by("-created_at")

        total = receipts.count()
        start = (page - 1) * per_page
        end = start + per_page
        receipts_page = receipts[start:end]

        results = [
            {
                "id": r.id,
                "transaction_number": r.transaction_number,
                "transaction_type": r.transaction_type,
                "transaction_type_display": r.get_transaction_type_display(),
                "amount": str(r.amount),
                "payment_mode_display": r.get_payment_mode_display(),
                "account_number": r.member_account.account_number,
                "balance_after": str(r.balance_after),
                "created_at": r.created_at.strftime("%b %d, %Y") if r.created_at else "",
            }
            for r in receipts_page
        ]

        return JsonResponse(
            {
                "success": True,
                "transactions": results,
                "total": total,
                "page": page,
                "has_more": end < total,
            }
        )
    except User.DoesNotExist:
        return JsonResponse({"success": False, "error": "Member not found"})


# ========================================
# Loan Views (Phase 3: LoanApplication + LoanAccount)
# ========================================

LOAN_LIST_STATUS_CHOICES = [
    ("pending", "Pending Approval"),
    ("rejected", "Rejected"),
    ("active", "Active"),
    ("closed", "Closed"),
    ("defaulted", "Defaulted"),
    ("written_off", "Written Off"),
]


@login_required
@admin_required
def loans_view(request):
    """Loans list: open applications (pending/rejected) + loan accounts."""
    page = request.GET.get("page", 1)
    query = request.GET.get("q", "").strip()
    loan_type = request.GET.get("type", "").strip()
    status = request.GET.get("status", "").strip()
    per_page = 25

    accounts_qs = LoanAccount.objects.select_related("application", "user", "disbursement_account", "created_by").all()
    open_apps_qs = (
        LoanApplication.objects.filter(status__in=["pending", "rejected"])
        .select_related("user", "created_by")
        .exclude(pk__in=LoanAccount.objects.values_list("application_id", flat=True))
    )

    if loan_type:
        accounts_qs = accounts_qs.filter(application__loan_type=loan_type)
        open_apps_qs = open_apps_qs.filter(loan_type=loan_type)

    if status:
        if status in ("active", "closed", "defaulted", "written_off"):
            accounts_qs = accounts_qs.filter(status=status)
            open_apps_qs = open_apps_qs.none()
        elif status in ("pending", "rejected"):
            accounts_qs = accounts_qs.none()
            open_apps_qs = open_apps_qs.filter(status=status)
        else:
            accounts_qs = accounts_qs.none()
            open_apps_qs = open_apps_qs.none()

    if query:
        accounts_qs = accounts_qs.filter(
            Q(loan_number__icontains=query)
            | Q(user__first_name__icontains=query)
            | Q(user__last_name__icontains=query)
            | Q(user__member_id__icontains=query)
            | Q(application__guarantor_name__icontains=query)
        )
        open_apps_qs = open_apps_qs.filter(
            Q(application_number__icontains=query)
            | Q(user__first_name__icontains=query)
            | Q(user__last_name__icontains=query)
            | Q(user__member_id__icontains=query)
            | Q(guarantor_name__icontains=query)
        )

    rows = [LoanListItem(account=a) for a in accounts_qs.order_by("-created_at")]
    rows.extend(LoanListItem(application=a) for a in open_apps_qs.order_by("-created_at"))
    rows.sort(key=lambda r: r.created_at, reverse=True)

    paginator = Paginator(rows, per_page)
    try:
        loans_page = paginator.page(page)
    except PageNotAnInteger:
        loans_page = paginator.page(1)
    except EmptyPage:
        loans_page = paginator.page(paginator.num_pages)

    _open_application_count = LoanApplication.objects.exclude(
        pk__in=LoanAccount.objects.values_list("application_id", flat=True)
    ).count()
    total_loans = LoanAccount.objects.count() + _open_application_count
    active_loans = LoanAccount.objects.filter(status="active").count()
    total_disbursed = LoanAccount.objects.filter(status="active").aggregate(total=Sum("principal_amount"))["total"] or 0
    total_outstanding = (
        LoanAccount.objects.filter(status="active").aggregate(total=Sum("outstanding_balance"))["total"] or 0
    )
    loan_label_map = dict(LoanApplication.LOAN_TYPE_CHOICES)
    configured_loan_types = list(
        LoanTypeConfiguration.objects.filter(is_active=True).order_by("display_order", "loan_type")
    )
    if configured_loan_types:
        loan_type_create_options = [
            {
                "value": cfg.loan_type,
                "label": loan_label_map.get(cfg.loan_type, cfg.loan_type.title()),
                "default_rate": str(cfg.interest_rate),
            }
            for cfg in configured_loan_types
        ]
    else:
        loan_type_create_options = [
            {"value": value, "label": label, "default_rate": ""} for value, label in LoanApplication.LOAN_TYPE_CHOICES
        ]
    penalty_per_day = SocietyConfiguration.objects.order_by("id").values_list(
        "late_payment_penalty_per_day", flat=True
    ).first() or Decimal("0.00")

    return render(
        request,
        "admin/loans.html",
        {
            "loans": loans_page,
            "search_query": query,
            "selected_type": loan_type,
            "selected_status": status,
            "loan_type_choices": LoanApplication.LOAN_TYPE_CHOICES,
            "status_choices": LOAN_LIST_STATUS_CHOICES,
            "page_obj": loans_page,
            "total_loans": total_loans,
            "active_loans": active_loans,
            "total_disbursed": total_disbursed,
            "total_outstanding": total_outstanding,
            "loan_type_create_options": loan_type_create_options,
            "late_penalty_per_day": penalty_per_day,
        },
    )


@login_required
@admin_required
def get_loan_view(request, loan_id):
    """Get loan application or loan account JSON for the modal. Use ?kind=application|account."""
    kind = request.GET.get("kind", "account").strip()
    try:
        if kind == "application":
            app = LoanApplication.objects.select_related("user", "approved_by", "created_by").get(id=loan_id)
            emi = app.calculate_emi()
            principal = app.principal_amount
            rate = app.interest_rate
            tenure = app.tenure_months
            if rate > 0:
                if app.interest_type == "flat":
                    total_interest = principal * rate * Decimal(str(tenure)) / Decimal("1200")
                    total_payable = (principal + total_interest).quantize(Decimal("0.01"))
                else:
                    total_payable = (emi * Decimal(str(tenure))).quantize(Decimal("0.01"))
            else:
                total_payable = principal
            return JsonResponse(
                {
                    "success": True,
                    "loan": {
                        "list_kind": "application",
                        "id": app.id,
                        "loan_number": app.application_number,
                        "loan_type": app.loan_type,
                        "loan_type_display": app.get_loan_type_display(),
                        "status": app.status,
                        "status_display": app.get_status_display(),
                        "member_name": app.user.display_name,
                        "member_id": app.user.member_id or "",
                        "user_id": app.user.id,
                        "principal_amount": str(app.principal_amount),
                        "interest_rate": str(app.interest_rate),
                        "interest_type": app.interest_type,
                        "interest_type_display": app.get_interest_type_display(),
                        "tenure_months": app.tenure_months,
                        "emi_amount": str(emi),
                        "total_payable": str(total_payable),
                        "total_paid": "0.00",
                        "outstanding_balance": str(app.principal_amount),
                        "overdue_amount": "0.00",
                        "completion_percentage": 0,
                        "application_date": app.application_date.strftime("%b %d, %Y") if app.application_date else "",
                        "approval_date": app.approval_date.strftime("%b %d, %Y") if app.approval_date else "",
                        "disbursement_date": "",
                        "first_emi_date": "",
                        "last_emi_date": "",
                        "closure_date": "",
                        "total_emis": app.tenure_months,
                        "emis_paid": 0,
                        "emis_overdue": 0,
                        "guarantor_name": app.guarantor_name or "",
                        "guarantor_member_id": app.guarantor_member_id or "",
                        "guarantor_relationship": app.guarantor_relationship or "",
                        "guarantor_contact": app.guarantor_contact or "",
                        "collateral_type": app.collateral_type or "",
                        "collateral_value": str(app.collateral_value) if app.collateral_value else "",
                        "collateral_description": app.collateral_description or "",
                        "disbursement_account_number": "",
                        "purpose": app.purpose or "",
                        "remarks": app.remarks or "",
                        "approved_by_name": app.approved_by.display_name if app.approved_by else "",
                        "created_by_name": app.created_by.display_name if app.created_by else "",
                        "created_at": app.created_at.strftime("%b %d, %Y") if app.created_at else "",
                        "repayments": [],
                    },
                }
            )

        loan = LoanAccount.objects.select_related(
            "application",
            "user",
            "disbursement_account",
            "created_by",
            "application__approved_by",
            "application__created_by",
        ).get(id=loan_id)
        app = loan.application
        g = loan.guarantors.order_by("id").first()
        repayments = loan.repayments.all().order_by("installment_number")
        repayments_data = [
            {
                "installment_number": r.installment_number,
                "due_date": r.due_date.strftime("%b %d, %Y") if r.due_date else "",
                "due_date_raw": r.due_date.strftime("%Y-%m-%d") if r.due_date else "",
                "paid_date": r.paid_date.strftime("%b %d, %Y") if r.paid_date else "",
                "amount_due": str(r.amount_due),
                "amount_paid": str(r.amount_paid),
                "principal_component": str(r.principal_component),
                "interest_component": str(r.interest_component),
                "penalty": str(r.penalty),
                "balance_after": str(r.balance_after),
                "payment_status": r.payment_status,
                "payment_status_display": r.get_payment_status_display(),
            }
            for r in repayments
        ]

        return JsonResponse(
            {
                "success": True,
                "loan": {
                    "list_kind": "account",
                    "id": loan.id,
                    "loan_number": loan.loan_number,
                    "loan_type": app.loan_type,
                    "loan_type_display": loan.get_loan_type_display(),
                    "status": loan.status,
                    "status_display": loan.get_status_display(),
                    "member_name": loan.user.display_name,
                    "member_id": loan.user.member_id or "",
                    "user_id": loan.user.id,
                    "principal_amount": str(loan.principal_amount),
                    "interest_rate": str(loan.interest_rate),
                    "interest_type": loan.interest_type,
                    "interest_type_display": loan.get_interest_type_display(),
                    "tenure_months": loan.tenure_months,
                    "emi_amount": str(loan.emi_amount),
                    "total_payable": str(loan.total_payable),
                    "total_paid": str(loan.total_paid),
                    "outstanding_balance": str(loan.outstanding_balance),
                    "overdue_amount": str(loan.overdue_amount),
                    "completion_percentage": loan.completion_percentage,
                    "application_date": app.application_date.strftime("%b %d, %Y") if app.application_date else "",
                    "approval_date": app.approval_date.strftime("%b %d, %Y") if app.approval_date else "",
                    "disbursement_date": loan.disbursement_date.strftime("%b %d, %Y") if loan.disbursement_date else "",
                    "first_emi_date": loan.first_emi_date.strftime("%b %d, %Y") if loan.first_emi_date else "",
                    "last_emi_date": loan.last_emi_date.strftime("%b %d, %Y") if loan.last_emi_date else "",
                    "closure_date": loan.closure_date.strftime("%b %d, %Y") if loan.closure_date else "",
                    "total_emis": loan.total_emis,
                    "emis_paid": loan.emis_paid,
                    "emis_overdue": loan.emis_overdue,
                    "guarantor_name": (g.name if g else (app.guarantor_name or "")),
                    "guarantor_member_id": (
                        (g.user.member_id if g and g.user else None) or app.guarantor_member_id or ""
                    ),
                    "guarantor_relationship": (g.relationship if g else app.guarantor_relationship) or "",
                    "guarantor_contact": (g.contact if g else app.guarantor_contact) or "",
                    "collateral_type": loan.collateral_type or app.collateral_type or "",
                    "collateral_value": str(loan.collateral_value or app.collateral_value or "")
                    if (loan.collateral_value or app.collateral_value)
                    else "",
                    "collateral_description": loan.collateral_description or app.collateral_description or "",
                    "disbursement_account_number": loan.disbursement_account.account_number
                    if loan.disbursement_account
                    else "",
                    "purpose": app.purpose or "",
                    "remarks": loan.remarks or app.remarks or "",
                    "approved_by_name": app.approved_by.display_name if app.approved_by else "",
                    "created_by_name": loan.created_by.display_name if loan.created_by else "",
                    "created_at": loan.created_at.strftime("%b %d, %Y") if loan.created_at else "",
                    "repayments": repayments_data,
                },
            }
        )
    except (LoanApplication.DoesNotExist, LoanAccount.DoesNotExist):
        return JsonResponse({"success": False, "error": "Loan not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def add_loan_view(request):
    """Create a new loan application (delegates to `services.loans.create_loan_application`)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})

    try:
        app = loan_service.create_loan_application(
            user_id=request.POST.get("user_id"),
            loan_type=request.POST.get("loan_type"),
            principal_amount=request.POST.get("principal_amount"),
            interest_rate=request.POST.get("interest_rate"),
            interest_type=request.POST.get("interest_type", "reducing"),
            tenure_months=request.POST.get("tenure_months"),
            purpose=request.POST.get("purpose", ""),
            remarks=request.POST.get("remarks", ""),
            guarantor_name=request.POST.get("guarantor_name", ""),
            guarantor_member_id=request.POST.get("guarantor_member_id", ""),
            guarantor_relationship=request.POST.get("guarantor_relationship", ""),
            guarantor_contact=request.POST.get("guarantor_contact", ""),
            collateral_type=request.POST.get("collateral_type", ""),
            collateral_value=request.POST.get("collateral_value") or None,
            collateral_description=request.POST.get("collateral_description", ""),
            processing_fee=request.POST.get("processing_fee") or Decimal("0"),
            disbursement_account_id=request.POST.get("disbursement_account_id") or None,
            actor=request.user,
            ip_address=_get_client_ip(request),
            audit_via="panel",
        )
    except ServiceError as e:
        return JsonResponse({"success": False, "error": e.message})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": True, "loan_id": app.id, "list_kind": "application"})


@login_required
@admin_required
def audit_logs_view(request):
    """Audit logs list view with filtering and pagination"""
    search_query = request.GET.get("q", "").strip()
    action_filter = request.GET.get("action", "").strip()
    entity_filter = request.GET.get("entity", "").strip()

    logs = AuditLog.objects.select_related("user").all()

    if search_query:
        logs = logs.filter(
            Q(description__icontains=search_query)
            | Q(user__first_name__icontains=search_query)
            | Q(user__last_name__icontains=search_query)
            | Q(user__username__icontains=search_query)
        )

    if action_filter:
        logs = logs.filter(action=action_filter)

    if entity_filter:
        logs = logs.filter(entity_type=entity_filter)

    # Pagination
    paginator = Paginator(logs, 50)
    page = request.GET.get("page")
    try:
        page_obj = paginator.page(page)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    context = {
        "logs": page_obj,
        "page_obj": page_obj,
        "search_query": search_query,
        "action_filter": action_filter,
        "entity_filter": entity_filter,
        "action_choices": AuditLog.ACTION_CHOICES,
        "entity_choices": AuditLog.ENTITY_CHOICES,
        "total_logs": paginator.count,
    }

    return render(request, "admin/audit_logs.html", context)


@login_required
@admin_required
def get_audit_log_view(request, log_id):
    """Get single audit log details"""
    try:
        log = AuditLog.objects.select_related("user").get(id=log_id)
        return JsonResponse(
            {
                "success": True,
                "log": {
                    "id": log.id,
                    "user_name": log.user.display_name if log.user else "System",
                    "user_username": log.user.username if log.user else "",
                    "action": log.action,
                    "action_display": log.get_action_display(),
                    "entity_type": log.entity_type,
                    "entity_type_display": log.get_entity_type_display(),
                    "entity_id": log.entity_id,
                    "description": log.description,
                    "ip_address": log.ip_address or "",
                    "created_at": log.created_at.strftime("%b %d, %Y %I:%M %p"),
                },
            }
        )
    except AuditLog.DoesNotExist:
        return JsonResponse({"success": False, "error": "Audit log not found"})


@login_required
@admin_required
def reports_view(request):
    """Financial reports page with profit calculation and portfolio analysis."""
    rd = report_dates_from_request(request)
    period = rd["period"]
    start_date = rd["start_date"]
    end_date = rd["end_date"]
    year = rd["year"]
    month = rd["month"]
    month_key = rd["month_key"]
    account_type_filter = rd["account_type_filter"]

    summary = get_financial_summary(start_date, end_date)
    account_type_cashflow = get_account_type_cashflow(months=6, account_type=account_type_filter or None)

    # Monthly breakdown for chart (revenue vs expense per month)
    monthly_data = []
    import calendar

    current = start_date
    while current <= end_date:
        month_start = current.replace(day=1)
        last_day = calendar.monthrange(current.year, current.month)[1]
        month_end = current.replace(day=last_day)
        month_summary = get_financial_summary(month_start, month_end)
        monthly_data.append(
            {
                "month": current.strftime("%b %Y"),
                "revenue": float(month_summary["gross_revenue"]),
                "expense": float(month_summary["total_expenses"]),
                "profit": float(month_summary["net_profit"]),
            }
        )
        # Move to next month
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1, day=1)
        else:
            current = current.replace(month=current.month + 1, day=1)

    # Fund balances
    fund_balances = FundAccount.objects.filter(is_deleted=False, is_active=True)

    # NPA summary
    npa_loans = [la for la in LoanAccount.objects.filter(status="active") if la.is_npa]
    npa_summary = {}
    for loan in npa_loans:
        cat = loan.npa_category or "substandard"
        if cat not in npa_summary:
            npa_summary[cat] = {"count": 0, "amount": Decimal("0.00")}
        npa_summary[cat]["count"] += 1
        npa_summary[cat]["amount"] += loan.outstanding_balance

    # Maturity alerts - accounts maturing in next 30/60/90 days
    today = date.today()
    maturing_30 = (
        MemberAccount.objects.filter(
            is_deleted=False, status="active", maturity_date__range=(today, today + timedelta(days=30))
        )
        .select_related("user")
        .order_by("maturity_date")
    )
    maturing_60 = MemberAccount.objects.filter(
        is_deleted=False, status="active", maturity_date__range=(today + timedelta(days=31), today + timedelta(days=60))
    ).select_related("user")
    maturing_90 = MemberAccount.objects.filter(
        is_deleted=False, status="active", maturity_date__range=(today + timedelta(days=61), today + timedelta(days=90))
    ).select_related("user")

    # Dividend summary (aggregated from ShareCapital lots)
    total_share_capital = ShareCapital.objects.filter(
        user__is_deleted=False,
        user__status="active",
        status="issued",
        redemption_date__isnull=True,
    ).aggregate(total=Sum("total_value"))["total"] or Decimal("0.00")
    eligible_dividend_members = (
        User.objects.filter(
            is_deleted=False,
            status="active",
            eligible_for_dividend=True,
            share_holdings__status="issued",
            share_holdings__redemption_date__isnull=True,
        )
        .distinct()
        .count()
    )

    report_fp = financial_period_overlapping(start_date, end_date)
    snapshot_pl = None
    snapshot_society = None
    if report_fp:
        snapshot_pl = ProfitAndLoss.objects.filter(financial_period=report_fp).first()
        snapshot_society = SocietyAccount.objects.filter(financial_period=report_fp).first()
    financial_periods = FinancialPeriod.objects.order_by("-start_date")[:80]

    context = {
        "summary": summary,
        "monthly_data_json": json.dumps(monthly_data),
        "fund_balances": fund_balances,
        "npa_summary": npa_summary,
        "npa_loans": npa_loans,
        "period": period,
        "year": year,
        "start_date": start_date,
        "end_date": end_date,
        "current_year": date.today().year,
        "maturing_30": maturing_30,
        "maturing_60_count": maturing_60.count(),
        "maturing_90_count": maturing_90.count(),
        "total_share_capital": total_share_capital,
        "eligible_dividend_members": eligible_dividend_members,
        "account_type_cashflow": account_type_cashflow,
        "selected_account_type": account_type_filter,
        "account_type_choices": MemberAccount.ACCOUNT_TYPE_CHOICES,
        "month": month,
        "month_key": month_key,
        "selected_quarter": rd["selected_quarter"],
        "fy_month_options": rd["fy_month_options"],
        "fy_quarters": rd["fy_quarters"],
        "working_fy": rd["working_fy"],
        "fy_window": rd["fy_window"],
        "using_session_fy": rd["using_session_fy"],
        "fy_year_label": rd["fy_year_label"],
        "report_financial_period": report_fp,
        "snapshot_pl": snapshot_pl,
        "snapshot_society": snapshot_society,
        "financial_periods": financial_periods,
    }
    return render(request, "admin/reports.html", context)


@login_required
@admin_required
def reports_save_financial_snapshot_view(request):
    """Persist P&L and society snapshot for a financial period (portal JSON)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"}, status=405)
    try:
        fp_id = request.POST.get("financial_period_id", "").strip()
        if fp_id.isdigit():
            fp = get_object_or_404(FinancialPeriod, pk=int(fp_id))
        else:
            fp = active_financial_period()
            if not fp:
                return JsonResponse(
                    {
                        "success": False,
                        "error": "No active financial period. Create one under Settings or pass financial_period_id.",
                    },
                    status=400,
                )

        pl_row, soc_row, pl_updated = persist_financial_snapshots(fp, calculated_by=request.user)
        log_action(
            request,
            "create",
            "system",
            None,
            f"Saved financial snapshot for {fp.label} (P&L updated={pl_updated}).",
        )
        return JsonResponse(
            {
                "success": True,
                "message": f"Snapshot saved for {fp.label}.",
                "profit_and_loss_updated": pl_updated,
                "net_surplus": str(soc_row.net_surplus),
                "total_income": str(pl_row.total_income),
            }
        )
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)


@login_required
@admin_required
def reports_sync_receivables_view(request):
    """Sync loan interest receivables from unpaid EMIs (portal JSON)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"}, status=405)
    try:
        from django.utils.dateparse import parse_date

        as_of = date.today()
        raw = request.POST.get("as_of_date", "").strip()
        if raw:
            parsed = parse_date(raw)
            if not parsed:
                return JsonResponse({"success": False, "error": "Invalid as_of_date; use YYYY-MM-DD"}, status=400)
            as_of = parsed

        fp = None
        fp_raw = request.POST.get("financial_period_id", "").strip()
        if fp_raw.isdigit():
            fp = get_object_or_404(FinancialPeriod, pk=int(fp_raw))

        n = sync_loan_interest_receivables(as_of_date=as_of, financial_period=fp)
        log_action(
            request,
            "update",
            "system",
            None,
            f"Synced loan interest receivables as of {as_of} ({n} installments).",
        )
        return JsonResponse(
            {
                "success": True,
                "message": f"Processed {n} unpaid installment(s) due on or before {as_of}.",
                "installments_considered": n,
            }
        )
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)


@login_required
@admin_required
def finance_fees_view(request):
    """Fee schedules and posted charges (read-only register)."""
    schedules = FeeSchedule.objects.all().order_by("-effective_date", "fee_type")
    q = request.GET.get("q", "").strip()
    charges = FeeCharge.objects.select_related("fee_schedule", "user", "loan_account").order_by("-created_at")
    if q:
        charges = charges.filter(
            Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q)
            | Q(user__member_id__icontains=q)
            | Q(fee_schedule__name__icontains=q)
        )
    page = request.GET.get("page", 1)
    per_page = 30
    paginator = Paginator(charges, per_page)
    try:
        charges_page = paginator.page(page)
    except PageNotAnInteger:
        charges_page = paginator.page(1)
    except EmptyPage:
        charges_page = paginator.page(paginator.num_pages)

    return render(
        request,
        "admin/finance_fees.html",
        {
            "schedules": schedules,
            "charges_page": charges_page,
            "search_query": q,
        },
    )


@login_required
@admin_required
def finance_receivables_view(request):
    """Interest receivable register."""
    qs = InterestReceivable.objects.select_related("loan_account", "financial_period", "loan_account__user").order_by(
        "-due_date", "-created_at"
    )
    st = request.GET.get("status", "").strip()
    if st:
        qs = qs.filter(status=st)
    page = request.GET.get("page", 1)
    per_page = 40
    paginator = Paginator(qs, per_page)
    try:
        page_obj = paginator.page(page)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    return render(
        request,
        "admin/finance_receivables.html",
        {"page_obj": page_obj, "status_filter": st},
    )


_VALID_SETTINGS_TABS = {"general", "financial-year", "products"}


def _settings_redirect(tab="general", next_path=None):
    """Redirect back to settings (preserving tab) or to a safe next path."""
    if next_path:
        return redirect(next_path)
    tab = tab if tab in _VALID_SETTINGS_TABS else "general"
    return redirect(f"/settings/?tab={tab}")


def _parse_decimal(raw, default=Decimal("0")):
    try:
        return Decimal((raw or "").strip() or "0")
    except Exception:
        return default


@login_required
@admin_required
def settings_view(request):
    """Settings hub: General profile, Financial Year, and Products (interest rates)."""
    from django.contrib import messages
    from django.utils.dateparse import parse_date

    requested_tab = request.GET.get("tab") or request.POST.get("tab") or "general"
    if requested_tab not in _VALID_SETTINGS_TABS:
        requested_tab = "general"

    if request.method == "POST":
        action = request.POST.get("action", "")

        # ---- General tab ----------------------------------------------------
        if action == "update_society_profile":
            society_name = request.POST.get("society_name", "").strip()
            penalty_raw = request.POST.get("late_payment_penalty_per_day", "").strip()
            remove_logo = request.POST.get("remove_logo") == "on"
            if not society_name:
                messages.error(request, "Society name is required.")
                return _settings_redirect("general")
            penalty_value = _parse_decimal(penalty_raw)
            if penalty_value < 0:
                messages.error(request, "Penalty per day cannot be negative.")
                return _settings_redirect("general")
            config, _ = SocietyConfiguration.objects.get_or_create(
                pk=1,
                defaults={
                    "society_name": society_name,
                    "late_payment_penalty_per_day": penalty_value,
                },
            )
            config.society_name = society_name
            config.late_payment_penalty_per_day = penalty_value
            logo = request.FILES.get("society_logo")
            if logo:
                config.society_logo = logo
            elif remove_logo and config.society_logo:
                config.society_logo.delete(save=False)
                config.society_logo = None
            config.save()
            log_action(request, "update", "system", config.id, "Updated society profile")
            messages.success(request, "General settings saved.")
            return _settings_redirect("general")

        # ---- Financial Year tab --------------------------------------------
        if action == "select_working_fy":
            pid = request.POST.get("working_period_id", "").strip()
            next_path = _safe_same_origin_redirect_path(request.POST.get("next", ""))
            if pid.isdigit():
                get_object_or_404(FinancialPeriod, pk=int(pid))
                set_working_financial_period(request, int(pid))
                fp = FinancialPeriod.objects.get(pk=int(pid))
                FinancialPeriod.objects.exclude(pk=fp.pk).update(is_active=False)
                fp.is_active = True
                fp.save(update_fields=["is_active"])
                messages.success(request, f"Working financial year set to {fp.label}.")
            else:
                set_working_financial_period(request, None)
                messages.success(request, "Working FY cleared. Using calendar Apr–Mar from today.")
            return _settings_redirect("financial-year", next_path=next_path)

        if action == "set_active":
            pid = request.POST.get("period_id", "")
            if pid.isdigit():
                fp = get_object_or_404(FinancialPeriod, pk=int(pid))
                FinancialPeriod.objects.exclude(pk=fp.pk).update(is_active=False)
                fp.is_active = True
                fp.save(update_fields=["is_active"])
                set_working_financial_period(request, fp.pk)
                messages.success(request, f"Active period is now {fp.label}.")
            return _settings_redirect("financial-year")

        if action == "delete_period":
            pid = request.POST.get("period_id", "")
            if pid.isdigit():
                fp = get_object_or_404(FinancialPeriod, pk=int(pid))
                if fp.is_active:
                    messages.error(request, "Cannot delete the active period. Activate another period first.")
                else:
                    label = fp.label
                    fp.delete()
                    messages.success(request, f"Deleted period {label}.")
            return _settings_redirect("financial-year")

        if action == "create_period":
            label = request.POST.get("label", "").strip()
            sd = parse_date(request.POST.get("start_date", ""))
            ed = parse_date(request.POST.get("end_date", ""))
            status_val = request.POST.get("status", "open")
            make_active = request.POST.get("is_active") == "on"
            if not label or not sd or not ed:
                messages.error(request, "Label, start date, and end date are required.")
                return _settings_redirect("financial-year")
            if ed < sd:
                messages.error(request, "End date must be on or after start date.")
                return _settings_redirect("financial-year")
            fp = FinancialPeriod.objects.create(
                label=label,
                start_date=sd,
                end_date=ed,
                status=status_val if status_val in dict(FinancialPeriod.STATUS_CHOICES) else "open",
                is_active=make_active,
                created_by=request.user,
            )
            if make_active:
                FinancialPeriod.objects.exclude(pk=fp.pk).update(is_active=False)
                set_working_financial_period(request, fp.pk)
            log_action(request, "create", "system", fp.id, f"Created financial period {fp.label}")
            messages.success(request, f"Created period {fp.label}.")
            return _settings_redirect("financial-year")

        # ---- Products tab --------------------------------------------------
        if action in {"create_account_type", "update_account_type", "delete_account_type"}:
            return _handle_product_action(
                request,
                model=AccountTypeConfiguration,
                name_field="account_type",
                action=action,
                tab="products",
            )

        if action in {"create_loan_type", "update_loan_type", "delete_loan_type"}:
            return _handle_product_action(
                request,
                model=LoanTypeConfiguration,
                name_field="loan_type",
                action=action,
                tab="products",
            )

        messages.error(request, "Unknown settings action.")
        return _settings_redirect(requested_tab)

    periods = FinancialPeriod.objects.order_by("-start_date")
    working = resolve_portal_financial_period(request)
    session_override = request.session.get(PORTAL_FY_SESSION_KEY)
    society_config = SocietyConfiguration.objects.order_by("id").first()
    account_type_configs = AccountTypeConfiguration.objects.order_by("display_order", "account_type")
    loan_type_configs = LoanTypeConfiguration.objects.order_by("display_order", "loan_type")
    return render(
        request,
        "admin/settings.html",
        {
            "periods": periods,
            "working_fy": working,
            "session_fy_override": session_override,
            "society_config": society_config,
            "account_type_configs": account_type_configs,
            "loan_type_configs": loan_type_configs,
            "active_tab": requested_tab,
        },
    )


def _handle_product_action(request, *, model, name_field, action, tab):
    """Shared CRUD handler for AccountTypeConfiguration / LoanTypeConfiguration."""
    from django.contrib import messages

    if action.startswith("delete"):
        item_id = request.POST.get("item_id", "")
        if item_id.isdigit():
            obj = get_object_or_404(model, pk=int(item_id))
            label = getattr(obj, name_field)
            obj.delete()
            messages.success(request, f"Removed product {label}.")
        return _settings_redirect(tab)

    name = (request.POST.get("name") or "").strip()
    rate_raw = (request.POST.get("interest_rate") or "").strip()
    is_active = request.POST.get("is_active") == "on"

    if not name:
        messages.error(request, "Product name is required.")
        return _settings_redirect(tab)
    if len(name) > 100:
        messages.error(request, "Product name must be 100 characters or fewer.")
        return _settings_redirect(tab)
    rate = _parse_decimal(rate_raw, default=Decimal("-1"))
    if rate < 0:
        messages.error(request, f"Invalid interest rate for {name}.")
        return _settings_redirect(tab)

    if action.startswith("create"):
        if model.objects.filter(**{name_field: name}).exists():
            messages.error(request, f"A product named “{name}” already exists.")
            return _settings_redirect(tab)
        last_order = model.objects.order_by("-display_order").values_list("display_order", flat=True).first() or 0
        model.objects.create(
            **{name_field: name},
            interest_rate=rate,
            is_active=is_active,
            display_order=last_order + 1,
        )
        messages.success(request, f"Added product {name}.")
        return _settings_redirect(tab)

    item_id = request.POST.get("item_id", "")
    if not item_id.isdigit():
        messages.error(request, "Missing product id.")
        return _settings_redirect(tab)
    obj = get_object_or_404(model, pk=int(item_id))
    if model.objects.filter(**{name_field: name}).exclude(pk=obj.pk).exists():
        messages.error(request, f"Another product is already named “{name}”.")
        return _settings_redirect(tab)
    setattr(obj, name_field, name)
    obj.interest_rate = rate
    obj.is_active = is_active
    obj.save()
    messages.success(request, f"Updated product {name}.")
    return _settings_redirect(tab)


@login_required
@admin_required
def database_view(request):
    """Database schema overview page for admins."""
    model_counts = [
        {
            "name": "User",
            "table": User._meta.db_table,
            "count": User.objects.count(),
            "purpose": "Members/Admins and profile data",
        },
        {
            "name": "MemberAccount",
            "table": MemberAccount._meta.db_table,
            "count": MemberAccount.objects.count(),
            "purpose": "Deposit/share/OD account ledgers",
        },
        {
            "name": "Transaction",
            "table": Transaction._meta.db_table,
            "count": Transaction.objects.count(),
            "purpose": "Posted transactions",
        },
        {
            "name": "Voucher",
            "table": Voucher._meta.db_table,
            "count": Voucher.objects.count(),
            "purpose": "Staged vouchers",
        },
        {
            "name": "VoucherEntry",
            "table": VoucherEntry._meta.db_table,
            "count": VoucherEntry.objects.count(),
            "purpose": "Line items under voucher",
        },
        {
            "name": "LoanApplication",
            "table": LoanApplication._meta.db_table,
            "count": LoanApplication.objects.count(),
            "purpose": "Loan applications",
        },
        {
            "name": "LoanAccount",
            "table": LoanAccount._meta.db_table,
            "count": LoanAccount.objects.count(),
            "purpose": "Active loan book",
        },
        {
            "name": "Guarantor",
            "table": Guarantor._meta.db_table,
            "count": Guarantor.objects.count(),
            "purpose": "Loan guarantors",
        },
        {
            "name": "LoanRepayment",
            "table": LoanRepayment._meta.db_table,
            "count": LoanRepayment.objects.count(),
            "purpose": "EMI schedule and payments",
        },
        {
            "name": "FundAccount",
            "table": FundAccount._meta.db_table,
            "count": FundAccount.objects.count(),
            "purpose": "Society funds",
        },
        {
            "name": "FundTransaction",
            "table": FundTransaction._meta.db_table,
            "count": FundTransaction.objects.count(),
            "purpose": "Fund ledger entries",
        },
        {
            "name": "AuditLog",
            "table": AuditLog._meta.db_table,
            "count": AuditLog.objects.count(),
            "purpose": "Audit trail",
        },
        {
            "name": "FinancialPeriod",
            "table": FinancialPeriod._meta.db_table,
            "count": FinancialPeriod.objects.count(),
            "purpose": "Fiscal years for snapshots",
        },
        {
            "name": "FeeSchedule",
            "table": FeeSchedule._meta.db_table,
            "count": FeeSchedule.objects.count(),
            "purpose": "Configurable fees",
        },
        {
            "name": "FeeCharge",
            "table": FeeCharge._meta.db_table,
            "count": FeeCharge.objects.count(),
            "purpose": "Posted fee instances",
        },
        {
            "name": "InterestReceivable",
            "table": InterestReceivable._meta.db_table,
            "count": InterestReceivable.objects.count(),
            "purpose": "Loan interest receivable",
        },
        {
            "name": "ProfitAndLoss",
            "table": ProfitAndLoss._meta.db_table,
            "count": ProfitAndLoss.objects.count(),
            "purpose": "Persisted P&L snapshots",
        },
        {
            "name": "SocietyAccount",
            "table": SocietyAccount._meta.db_table,
            "count": SocietyAccount.objects.count(),
            "purpose": "Society balance sheet snapshots",
        },
    ]

    relationships = [
        {"from": "User", "to": "MemberAccount", "label": "1:N"},
        {"from": "User", "to": "Transaction", "label": "1:N"},
        {"from": "MemberAccount", "to": "Transaction", "label": "1:N"},
        {"from": "User", "to": "Voucher", "label": "1:N"},
        {"from": "Voucher", "to": "VoucherEntry", "label": "1:N"},
        {"from": "MemberAccount", "to": "VoucherEntry", "label": "1:N"},
        {"from": "User", "to": "LoanApplication", "label": "1:N"},
        {"from": "User", "to": "LoanAccount", "label": "1:N"},
        {"from": "LoanApplication", "to": "LoanAccount", "label": "1:1"},
        {"from": "LoanAccount", "to": "Guarantor", "label": "1:N"},
        {"from": "LoanAccount", "to": "LoanRepayment", "label": "1:N"},
        {"from": "Transaction", "to": "LoanRepayment", "label": "1:N"},
        {"from": "FundAccount", "to": "FundTransaction", "label": "1:N"},
        {"from": "User", "to": "AuditLog", "label": "1:N"},
    ]

    mermaid_lines = [
        "graph LR",
        'U["User"]',
        'MA["MemberAccount"]',
        'R["Transaction"]',
        'V["Voucher"]',
        'VE["VoucherEntry"]',
        'LA1["LoanApplication"]',
        'LA2["LoanAccount"]',
        'G["Guarantor"]',
        'LR["LoanRepayment"]',
        'F["FundAccount"]',
        'FT["FundTransaction"]',
        'AL["AuditLog"]',
        "U -->|1:N| MA",
        "U -->|1:N| R",
        "MA -->|1:N| R",
        "U -->|1:N| V",
        "V -->|1:N| VE",
        "MA -->|1:N| VE",
        "U -->|1:N| LA1",
        "U -->|1:N| LA2",
        "LA1 -->|1:1| LA2",
        "LA2 -->|1:N| G",
        "LA2 -->|1:N| LR",
        "R -->|1:N| LR",
        "F -->|1:N| FT",
        "U -->|1:N| AL",
    ]

    return render(
        request,
        "admin/database.html",
        {
            "model_counts": model_counts,
            "relationships": relationships,
            "mermaid_graph": "\n".join(mermaid_lines),
        },
    )


@login_required
@admin_required
def distribute_profit_view(request):
    """Distribute annual profit to funds using allocation rules.

    Uses the **locked** :class:`~accounts.models.ProfitAndLoss` ``net_surplus`` for the
    working financial period (or ``financial_period_id`` when posted). Rejects when
    no snapshot exists or P&L is not locked (Phase A6).
    """
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})
    try:
        from accounts.utils import apply_fund_allocations

        fp = None
        fp_raw = request.POST.get("financial_period_id", "").strip()
        if fp_raw.isdigit():
            fp = get_object_or_404(FinancialPeriod, pk=int(fp_raw))
        else:
            fw = effective_fy_window(request)
            fp = financial_period_overlapping(fw.start_date, fw.end_date)

        if not fp:
            return JsonResponse(
                {
                    "success": False,
                    "error": "No financial period matches the working FY. Open or select a financial year first.",
                }
            )

        pl = ProfitAndLoss.objects.filter(financial_period=fp).first()
        if not pl:
            return JsonResponse(
                {
                    "success": False,
                    "error": "No P&L snapshot for this financial period. Save a financial snapshot first.",
                }
            )
        if not pl.is_locked:
            return JsonResponse(
                {
                    "success": False,
                    "error": "P&L is not locked for this period. Close the financial year before distributing profit.",
                }
            )

        net_surplus = pl.net_surplus
        if net_surplus <= 0:
            return JsonResponse(
                {
                    "success": False,
                    "error": f"No distributable surplus for {fp.label}. Locked net surplus: ₹{net_surplus}.",
                }
            )

        transactions = apply_fund_allocations(
            request,
            "annual_profit",
            net_surplus,
            f"Annual profit distribution for {fp.label} (locked P&L net surplus)",
        )

        if not transactions:
            return JsonResponse(
                {
                    "success": False,
                    "error": "No allocation rules configured for 'Annual Profit'. Add rules in Allocation Rules page.",
                }
            )

        total_allocated = sum(t.amount for t in transactions)
        log_action(
            request,
            "create",
            "fund",
            None,
            f"Distributed locked P&L surplus ₹{net_surplus} for {fp.label}. "
            f"Allocated ₹{total_allocated} across {len(transactions)} funds.",
        )

        return JsonResponse(
            {
                "success": True,
                "message": f"Distributed ₹{total_allocated} across {len(transactions)} funds (from locked P&L).",
            }
        )
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def distribute_dividend_view(request):
    """Calculate and distribute dividends to eligible members based on share capital."""
    if request.method == "POST":
        try:
            fw = effective_fy_window(request)
            y0, y1 = fw.start_date.year, fw.end_date.year
            dividend_rate = request.POST.get("dividend_rate")

            if not dividend_rate:
                return JsonResponse({"success": False, "error": "Dividend rate is required"})

            dividend_rate = Decimal(dividend_rate)
            if dividend_rate <= 0 or dividend_rate > 100:
                return JsonResponse({"success": False, "error": "Dividend rate must be between 0.01 and 100"})

            eligible_members = (
                User.objects.filter(
                    is_deleted=False,
                    status="active",
                    eligible_for_dividend=True,
                )
                .annotate(
                    share_capital_total=Coalesce(
                        Sum(
                            "share_holdings__total_value",
                            filter=Q(
                                share_holdings__status="issued",
                                share_holdings__redemption_date__isnull=True,
                            ),
                        ),
                        Value(0),
                        output_field=DecimalField(max_digits=14, decimal_places=2),
                    )
                )
                .filter(share_capital_total__gt=0)
            )

            if not eligible_members.exists():
                return JsonResponse({"success": False, "error": "No eligible members with share capital found"})

            total_distributed = Decimal("0.00")
            member_count = 0

            with transaction.atomic():
                for member in eligible_members:
                    dividend_amount = (member.share_capital_total * dividend_rate / Decimal("100")).quantize(
                        Decimal("0.01")
                    )

                    if dividend_amount <= 0:
                        continue

                    # Update member's dividend payable
                    User.objects.filter(id=member.id).update(
                        dividend_payable_balance=F("dividend_payable_balance") + dividend_amount,
                        last_dividend_paid_date=date.today(),
                    )

                    total_distributed += dividend_amount
                    member_count += 1

                # Deduct from Dividend Fund if it exists
                dividend_fund = FundAccount.objects.filter(
                    fund_type="dividend", is_deleted=False, is_active=True
                ).first()

                if dividend_fund:
                    new_balance = dividend_fund.balance - total_distributed
                    FundAccount.objects.filter(id=dividend_fund.id).update(balance=F("balance") - total_distributed)
                    FundTransaction.objects.create(
                        fund=dividend_fund,
                        transaction_type="debit",
                        amount=total_distributed,
                        description=f"Dividend distribution {getattr(fw, 'label', 'FY')} ({y0}-{y1}) @ {dividend_rate}%",
                        payment_mode="internal",
                        balance_after=new_balance,
                        trigger_event="manual",
                        created_by=request.user,
                    )

                log_action(
                    request,
                    "create",
                    "fund",
                    None,
                    f"Distributed dividend @ {dividend_rate}% for {getattr(fw, 'label', 'FY')} ({y0}-{y1}). "
                    f"₹{total_distributed} to {member_count} members.",
                )

            return JsonResponse(
                {
                    "success": True,
                    "message": f"Distributed ₹{total_distributed} as dividend @ {dividend_rate}% to {member_count} members.",
                }
            )
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def post_interest_view(request):
    """Calculate and post accrued interest to eligible deposit accounts.

    Phase B1: delegates to ``accounts.services.interest_engine.accrue_deposits``.
    Behaviour: same set of accounts (CD/FD/RD/Sukanya/Suputra — `share`+`od`
    excluded), same per-account window
    ``[last_interest_calc_date or opening_date, today]``, same audit summary.
    """
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})
    try:
        fp = resolve_portal_financial_period(request) or active_financial_period()
        result = interest_engine.accrue_deposits(
            as_of=date.today(),
            financial_period=fp,
            actor=request.user,
            ip_address=_get_client_ip(request),
            audit_via="panel",
        )

        log_action(
            request,
            "create",
            "account",
            None,
            (
                f"Posted interest for period ending {result['period_end']}. "
                f"\u20b9{result['total_posted']} across {result['accounts_updated']} accounts."
            ),
        )

        return JsonResponse(
            {
                "success": True,
                "message": (
                    f"Posted \u20b9{result['total_posted']} interest to "
                    f"{result['accounts_updated']} accounts."
                ),
            }
        )
    except ServiceError as exc:
        return JsonResponse({"success": False, "error": str(exc)}, status=400)
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def export_report_view(request):
    """Export financial report data as CSV."""
    import calendar
    import csv

    from accounts.utils import get_financial_summary

    rd = report_dates_from_request(request)
    start_date = rd["start_date"]
    end_date = rd["end_date"]
    period = rd["period"]
    if period == "month":
        filename = f"financial_report_{start_date.year}_{start_date.month:02d}.csv"
    elif period == "quarter":
        filename = f"financial_report_Q{rd['selected_quarter']}_{start_date.year}.csv"
    else:
        safe_lbl = getattr(rd["fy_window"], "label", "FY").replace(" ", "_")
        filename = f"financial_report_{safe_lbl}.csv"

    summary = get_financial_summary(start_date, end_date)

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)

    # Header
    writer.writerow(["Financial Report"])
    writer.writerow(["Period", f"{start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}"])
    writer.writerow([])

    # Revenue
    writer.writerow(["REVENUE"])
    writer.writerow(["Loan Interest Income", float(summary["loan_interest_income"])])
    writer.writerow(["Penalty Income", float(summary["penalty_income"])])
    writer.writerow(["Processing Fees", float(summary["processing_fees"])])
    writer.writerow(["Gross Revenue", float(summary["gross_revenue"])])
    writer.writerow([])

    # Expenses
    writer.writerow(["EXPENSES"])
    writer.writerow(["Deposit Interest Expense", float(summary["deposit_interest_expense"])])
    writer.writerow(["Monthly Interest Liability", float(summary["monthly_interest_liability"])])
    writer.writerow(["Total Expenses", float(summary["total_expenses"])])
    writer.writerow([])

    # Profit
    writer.writerow(["NET PROFIT", float(summary["net_profit"])])
    writer.writerow([])

    # Loan Portfolio
    writer.writerow(["LOAN PORTFOLIO"])
    writer.writerow(["Active Loans", summary["active_loans_count"]])
    writer.writerow(["Total Disbursed (Period)", float(summary["total_disbursed"])])
    writer.writerow(["Total Outstanding", float(summary["total_loans_outstanding"])])
    writer.writerow(["Total Overdue", float(summary["total_overdue"])])
    writer.writerow(["NPA Count", summary["npa_count"]])
    writer.writerow(["Recovery Rate (%)", summary["recovery_rate"]])
    writer.writerow([])

    # Deposit Portfolio
    writer.writerow(["DEPOSIT PORTFOLIO"])
    writer.writerow(["Account Type", "Balance", "Avg Rate (%)", "Count"])
    for type_name, data in summary["deposit_by_type"].items():
        writer.writerow([type_name, float(data["balance"]), float(data["avg_rate"]), data["count"]])
    writer.writerow([])

    # Monthly Breakdown
    writer.writerow(["MONTHLY BREAKDOWN"])
    writer.writerow(["Month", "Revenue", "Expense", "Net Profit"])
    current = start_date
    while current <= end_date:
        month_start = current.replace(day=1)
        last_day = calendar.monthrange(current.year, current.month)[1]
        month_end = current.replace(day=last_day)
        month_summary = get_financial_summary(month_start, month_end)
        writer.writerow(
            [
                current.strftime("%b %Y"),
                float(month_summary["gross_revenue"]),
                float(month_summary["total_expenses"]),
                float(month_summary["net_profit"]),
            ]
        )
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1, day=1)
        else:
            current = current.replace(month=current.month + 1, day=1)

    log_action(request, "export", "system", None, f"Exported financial report for {start_date} to {end_date}")
    return response


@login_required
@admin_required
def calculator_view(request):
    """Banking calculator page"""
    return render(request, "admin/calculator.html")


@login_required
@admin_required
def export_members_view(request):
    """Export members list to Excel"""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"

    headers = [
        "Member ID",
        "Username",
        "Name",
        "Email",
        "Phone",
        "Status",
        "Member Type",
        "Date of Joining",
        "Date of Birth",
        "Gender",
        "Occupation",
        "KYC Status",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    members = User.objects.filter(is_deleted=False, role="member").order_by("member_id")
    for m in members:
        ws.append(
            [
                m.member_id or "",
                m.username,
                m.display_name,
                m.email or "",
                m.mobile_primary or "",
                m.get_status_display(),
                m.get_member_type_display(),
                m.date_of_joining.strftime("%Y-%m-%d") if m.date_of_joining else "",
                m.date_of_birth.strftime("%Y-%m-%d") if m.date_of_birth else "",
                m.get_gender_display() if m.gender else "",
                m.occupation or "",
                (user_kyc(m).get_kyc_status_display() if user_kyc(m) else ""),
            ]
        )

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

    log_action(request, "export", "member", None, "Exported members list to Excel")

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="members.xlsx"'
    wb.save(response)
    return response


@login_required
@admin_required
def export_accounts_view(request):
    """Export accounts list to Excel"""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts"

    headers = [
        "Account Number",
        "Member Name",
        "Member ID",
        "Account Type",
        "Status",
        "Balance",
        "Interest Rate (%)",
        "Principal Amount",
        "Opening Date",
        "Maturity Date",
        "Tenure (Months)",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    accounts = MemberAccount.objects.filter(is_deleted=False).select_related("user").order_by("account_number")
    for a in accounts:
        ws.append(
            [
                a.account_number,
                a.user.display_name,
                a.user.member_id or "",
                a.get_account_type_display(),
                a.get_status_display(),
                float(a.balance),
                float(a.interest_rate),
                float(a.principal_amount),
                a.opening_date.strftime("%Y-%m-%d") if a.opening_date else "",
                a.maturity_date.strftime("%Y-%m-%d") if a.maturity_date else "",
                a.tenure_months or "",
            ]
        )

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

    log_action(request, "export", "account", None, "Exported accounts list to Excel")

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="accounts.xlsx"'
    wb.save(response)
    return response


@login_required
@admin_required
def export_transactions_view(request):
    """Export transactions list to Excel."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receipts"

    headers = [
        "Receipt Number",
        "Member Name",
        "Member ID",
        "Account Number",
        "Transaction Type",
        "Amount",
        "Payment Mode",
        "Reference Number",
        "Balance After",
        "Created By",
        "Date",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    receipts = (
        Transaction.objects.select_related("user", "member_account", "created_by").all().order_by("-created_at")[:5000]
    )
    for r in receipts:
        ws.append(
            [
                r.transaction_number,
                r.user.display_name,
                r.user.member_id or "",
                r.member_account.account_number if r.member_account else "",
                r.get_transaction_type_display(),
                float(r.amount),
                r.get_payment_mode_display(),
                r.reference_number or "",
                float(r.balance_after),
                r.created_by.display_name if r.created_by else "",
                r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
            ]
        )

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

    log_action(request, "export", "transaction", None, "Exported receipts list to Excel")

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="receipts.xlsx"'
    wb.save(response)
    return response


@login_required
@admin_required
def export_loans_view(request):
    """Export loans list to Excel"""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loans"

    headers = [
        "Loan Number",
        "Member Name",
        "Member ID",
        "Loan Type",
        "Status",
        "Principal Amount",
        "Interest Rate (%)",
        "Interest Type",
        "Tenure (Months)",
        "EMI Amount",
        "Total Payable",
        "Total Paid",
        "Outstanding Balance",
        "Application Date",
        "Approval Date",
        "Disbursement Date",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for ac in LoanAccount.objects.select_related("application", "user").order_by("-created_at"):
        app = ac.application
        ws.append(
            [
                ac.loan_number,
                ac.user.display_name,
                ac.user.member_id or "",
                ac.get_loan_type_display(),
                ac.get_status_display(),
                float(ac.principal_amount),
                float(ac.interest_rate),
                ac.get_interest_type_display(),
                ac.tenure_months,
                float(ac.emi_amount),
                float(ac.total_payable),
                float(ac.total_paid),
                float(ac.outstanding_balance),
                app.application_date.strftime("%Y-%m-%d") if app.application_date else "",
                app.approval_date.strftime("%Y-%m-%d") if app.approval_date else "",
                ac.disbursement_date.strftime("%Y-%m-%d") if ac.disbursement_date else "",
            ]
        )
    open_apps = (
        LoanApplication.objects.exclude(pk__in=LoanAccount.objects.values_list("application_id", flat=True))
        .select_related("user")
        .order_by("-created_at")
    )
    for app in open_apps:
        emi = app.calculate_emi()
        if app.interest_rate > 0:
            if app.interest_type == "flat":
                total_pay = float(
                    (
                        app.principal_amount
                        + app.principal_amount * app.interest_rate * Decimal(str(app.tenure_months)) / Decimal("1200")
                    ).quantize(Decimal("0.01"))
                )
            else:
                total_pay = float((emi * Decimal(str(app.tenure_months))).quantize(Decimal("0.01")))
        else:
            total_pay = float(app.principal_amount)
        ws.append(
            [
                app.application_number,
                app.user.display_name,
                app.user.member_id or "",
                app.get_loan_type_display(),
                app.get_status_display(),
                float(app.principal_amount),
                float(app.interest_rate),
                app.get_interest_type_display(),
                app.tenure_months,
                float(emi),
                total_pay,
                0.0,
                float(app.principal_amount),
                app.application_date.strftime("%Y-%m-%d") if app.application_date else "",
                app.approval_date.strftime("%Y-%m-%d") if app.approval_date else "",
                "",
            ]
        )

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 40)

    log_action(request, "export", "loan", None, "Exported loans list to Excel")

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="loans.xlsx"'
    wb.save(response)
    return response


@login_required
@admin_required
def approve_loan_view(request, loan_id):
    """Approve a pending LoanApplication (delegates to `services.loans.approve_loan_application`)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})

    raw_acct = request.POST.get("disbursement_account_id") or request.POST.get("disbursement_account")
    disbursement_account_id = None
    if raw_acct not in (None, ""):
        try:
            disbursement_account_id = int(raw_acct)
        except (TypeError, ValueError):
            return JsonResponse({"success": False, "error": "Invalid disbursement account ID."}, status=400)

    processing_fee = request.POST.get("processing_fee")
    if processing_fee is not None and str(processing_fee).strip() == "":
        processing_fee = None

    try:
        acct = loan_service.approve_loan_application(
            application_id=loan_id,
            disbursement_date=request.POST.get("disbursement_date"),
            disbursement_account_id=disbursement_account_id,
            processing_fee=processing_fee,
            actor=request.user,
            ip_address=_get_client_ip(request),
            audit_via="panel",
        )
    except ServiceError as e:
        return JsonResponse({"success": False, "error": e.message})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": True, "loan_account_id": acct.id})


@login_required
@admin_required
def record_emi_payment_view(request, loan_id):
    """Record an EMI payment (delegates to `services.loans.record_emi_payment`)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"})

    try:
        loan_service.record_emi_payment(
            loan_account_id=loan_id,
            installment_number=request.POST.get("installment_number"),
            amount_paid=request.POST.get("amount_paid"),
            payment_mode=request.POST.get("payment_mode", "cash"),
            actor=request.user,
            ip_address=_get_client_ip(request),
            audit_via="panel",
        )
    except ServiceError as e:
        return JsonResponse({"success": False, "error": e.message})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": True})


# ========================================
# Fund Management Views
# ========================================


@login_required
@admin_required
def funds_view(request):
    """Funds list view with pagination and search"""
    page = request.GET.get("page", 1)
    query = request.GET.get("q", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    per_page = 25

    # Base queryset
    funds = FundAccount.objects.filter(is_deleted=False).order_by("name")

    # Apply search filter
    if query:
        funds = funds.filter(Q(name__icontains=query) | Q(account_number__icontains=query))
    if date_from:
        funds = funds.filter(created_at__date__gte=date_from)
    if date_to:
        funds = funds.filter(created_at__date__lte=date_to)

    # Paginate
    paginator = Paginator(funds, per_page)
    try:
        funds_page = paginator.page(page)
    except PageNotAnInteger:
        funds_page = paginator.page(1)
    except EmptyPage:
        funds_page = paginator.page(paginator.num_pages)

    # Summary stats
    total_funds = FundAccount.objects.filter(is_deleted=False).count()
    total_balance = FundAccount.objects.filter(is_deleted=False).aggregate(total=Sum("balance"))["total"] or 0
    pending_vouchers = Voucher.objects.filter(status="pending").select_related("user").order_by("-created_at")[:50]
    active_funds = FundAccount.objects.filter(is_deleted=False, is_active=True).order_by("name")

    return render(
        request,
        "admin/funds.html",
        {
            "funds": funds_page,
            "search_query": query,
            "page_obj": funds_page,
            "total_funds": total_funds,
            "total_balance": total_balance,
            "date_from": date_from,
            "date_to": date_to,
            "pending_vouchers": pending_vouchers,
            "active_funds": active_funds,
        },
    )


@login_required
@admin_required
def add_fund_view(request):
    """Create a new fund account"""
    if request.method == "POST":
        try:
            name = request.POST.get("name", "").strip()
            fund_type = request.POST.get("fund_type")
            description = request.POST.get("description", "").strip()

            # Validate required fields
            if not name:
                return JsonResponse({"success": False, "error": "Fund name is required"})
            if not fund_type:
                return JsonResponse({"success": False, "error": "Fund type is required"})

            # Auto-generate account number: FUND-YEAR-SEQUENCE
            max_retries = 5
            for attempt in range(max_retries):
                try:
                    with transaction.atomic():
                        year = date.today().year
                        last_fund = (
                            FundAccount.objects.select_for_update()
                            .filter(account_number__startswith=f"FUND-{year}-")
                            .order_by("-account_number")
                            .first()
                        )
                        if last_fund:
                            try:
                                last_seq = int(last_fund.account_number.split("-")[-1])
                                next_seq = last_seq + 1
                            except (ValueError, IndexError):
                                next_seq = 1
                        else:
                            next_seq = 1

                        account_number = f"FUND-{year}-{next_seq:05d}"

                        fund = FundAccount.objects.create(
                            name=name,
                            fund_type=fund_type,
                            account_number=account_number,
                            description=description or None,
                            created_by=request.user,
                        )

                    log_action(
                        request,
                        "create",
                        "fund",
                        fund.id,
                        f"Created fund {name} ({account_number})",
                    )
                    return JsonResponse({"success": True, "fund_id": fund.id})
                except IntegrityError:
                    if attempt == max_retries - 1:
                        raise
                    continue

        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def get_fund_view(request, fund_id):
    """Get fund details with recent transactions and allocation rules"""
    try:
        fund = FundAccount.objects.get(id=fund_id, is_deleted=False)

        # Get recent transactions
        transactions = fund.transactions.all().order_by("-created_at")[:10]
        transactions_data = [
            {
                "id": t.id,
                "transaction_type": t.transaction_type,
                "transaction_type_display": t.get_transaction_type_display(),
                "amount": str(t.amount),
                "description": t.description or "",
                "payment_mode_display": t.get_payment_mode_display(),
                "balance_after": str(t.balance_after),
                "created_at": t.created_at.strftime("%b %d, %Y") if t.created_at else "",
            }
            for t in transactions
        ]

        # Get allocation rules
        rules = fund.allocation_rules.filter(is_active=True).order_by("trigger_event")
        rules_data = [
            {
                "id": r.id,
                "trigger_event": r.trigger_event,
                "trigger_event_display": r.get_trigger_event_display(),
                "allocation_type": r.allocation_type,
                "allocation_type_display": r.get_allocation_type_display(),
                "amount": str(r.amount) if r.amount else "",
                "percentage": str(r.percentage) if r.percentage else "",
            }
            for r in rules
        ]

        return JsonResponse(
            {
                "success": True,
                "fund": {
                    "id": fund.id,
                    "name": fund.name,
                    "fund_type": fund.fund_type,
                    "fund_type_display": fund.get_fund_type_display(),
                    "account_number": fund.account_number,
                    "balance": str(fund.balance),
                    "description": fund.description or "",
                    "is_active": fund.is_active,
                    "created_at": fund.created_at.strftime("%b %d, %Y") if fund.created_at else "",
                    "transactions": transactions_data,
                    "rules": rules_data,
                },
            }
        )
    except FundAccount.DoesNotExist:
        return JsonResponse({"success": False, "error": "Fund not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def edit_fund_view(request, fund_id):
    """Update fund details"""
    if request.method == "POST":
        try:
            fund = FundAccount.objects.get(id=fund_id, is_deleted=False)

            name = request.POST.get("name", "").strip()
            fund_type = request.POST.get("fund_type")
            description = request.POST.get("description", "").strip()
            is_active = request.POST.get("is_active") == "true"

            if not name:
                return JsonResponse({"success": False, "error": "Fund name is required"})
            if not fund_type:
                return JsonResponse({"success": False, "error": "Fund type is required"})

            fund.name = name
            fund.fund_type = fund_type
            fund.description = description or None
            fund.is_active = is_active
            fund.save()

            log_action(
                request,
                "update",
                "fund",
                fund.id,
                f"Updated fund {fund.name} ({fund.account_number})",
            )
            return JsonResponse({"success": True})
        except FundAccount.DoesNotExist:
            return JsonResponse({"success": False, "error": "Fund not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def delete_fund_view(request, fund_id):
    """Soft-delete a fund account (only if balance is zero)"""
    if request.method == "POST":
        try:
            fund = FundAccount.objects.get(id=fund_id, is_deleted=False)

            # Protect funds with non-zero balance
            if fund.balance > 0:
                return JsonResponse(
                    {
                        "success": False,
                        "error": f"Cannot delete fund with balance ₹{fund.balance}. Please transfer funds first.",
                    }
                )

            fund.is_deleted = True
            fund.deleted_at = timezone.now()
            fund.is_active = False
            fund.save()

            log_action(
                request,
                "delete",
                "fund",
                fund.id,
                f"Deleted fund {fund.name} ({fund.account_number})",
            )
            return JsonResponse({"success": True})
        except FundAccount.DoesNotExist:
            return JsonResponse({"success": False, "error": "Fund not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def add_fund_transaction_view(request, fund_id):
    """Add a manual fund transaction"""
    if request.method == "POST":
        try:
            transaction_type = request.POST.get("transaction_type")
            amount = request.POST.get("amount")
            description = request.POST.get("description", "").strip()
            payment_mode = request.POST.get("payment_mode", "cash")
            reference_number = request.POST.get("reference_number", "").strip()

            # Validate required fields
            if not transaction_type:
                return JsonResponse({"success": False, "error": "Transaction type is required"})
            if not amount:
                return JsonResponse({"success": False, "error": "Amount is required"})
            if not description:
                return JsonResponse({"success": False, "error": "Description is required"})

            amount_decimal = Decimal(str(amount))
            if amount_decimal <= 0:
                return JsonResponse({"success": False, "error": "Amount must be greater than zero"})

            with transaction.atomic():
                # Lock the fund account
                fund = FundAccount.objects.select_for_update().get(id=fund_id, is_deleted=False)

                # Calculate new balance
                if transaction_type == "credit":
                    new_balance = fund.balance + amount_decimal
                else:
                    # Check sufficient balance for debit
                    if amount_decimal > fund.balance:
                        return JsonResponse(
                            {
                                "success": False,
                                "error": f"Insufficient balance. Fund has ₹{fund.balance} but ₹{amount_decimal} was requested.",
                            }
                        )
                    new_balance = fund.balance - amount_decimal

                # Update balance using F() expression
                if transaction_type == "credit":
                    FundAccount.objects.filter(id=fund.id).update(balance=F("balance") + amount_decimal)
                else:
                    FundAccount.objects.filter(id=fund.id).update(balance=F("balance") - amount_decimal)

                # Create transaction record
                fund_txn = FundTransaction.objects.create(
                    fund=fund,
                    transaction_type=transaction_type,
                    amount=amount_decimal,
                    description=description,
                    payment_mode=payment_mode,
                    reference_number=reference_number or None,
                    balance_after=new_balance,
                    trigger_event="manual",
                    created_by=request.user,
                )

            log_action(
                request,
                "create",
                "fund",
                fund.id,
                f"Manual {transaction_type} of ₹{amount_decimal} to {fund.name}: {description}",
            )
            return JsonResponse({"success": True, "transaction_id": fund_txn.id})
        except FundAccount.DoesNotExist:
            return JsonResponse({"success": False, "error": "Fund not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def fund_transactions_view(request, fund_id):
    """Get paginated transaction history for a fund"""
    try:
        fund = FundAccount.objects.get(id=fund_id, is_deleted=False)
        page = int(request.GET.get("page", 1))
        per_page = 25

        transactions = fund.transactions.all().order_by("-created_at")
        total = transactions.count()
        start = (page - 1) * per_page
        end = start + per_page
        page_transactions = transactions[start:end]

        results = [
            {
                "id": t.id,
                "transaction_type": t.transaction_type,
                "transaction_type_display": t.get_transaction_type_display(),
                "amount": str(t.amount),
                "description": t.description or "",
                "payment_mode_display": t.get_payment_mode_display(),
                "reference_number": t.reference_number or "",
                "balance_after": str(t.balance_after),
                "trigger_event_display": t.get_trigger_event_display() if t.trigger_event else "",
                "source_member": t.source_member.display_name if t.source_member else "",
                "created_by": t.created_by.display_name if t.created_by else "",
                "created_at": t.created_at.strftime("%b %d, %Y %I:%M %p") if t.created_at else "",
            }
            for t in page_transactions
        ]

        return JsonResponse(
            {
                "success": True,
                "fund_name": fund.name,
                "fund_account_number": fund.account_number,
                "transactions": results,
                "total": total,
                "page": page,
                "has_more": end < total,
            }
        )
    except FundAccount.DoesNotExist:
        return JsonResponse({"success": False, "error": "Fund not found"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)})


@login_required
@admin_required
def allocation_rules_view(request):
    """Allocation rules list view with pagination"""
    page = request.GET.get("page", 1)
    per_page = 25

    # Base queryset
    rules = FundAllocationRule.objects.select_related("fund", "created_by").all().order_by("trigger_event", "fund")

    # Paginate
    paginator = Paginator(rules, per_page)
    try:
        rules_page = paginator.page(page)
    except PageNotAnInteger:
        rules_page = paginator.page(1)
    except EmptyPage:
        rules_page = paginator.page(paginator.num_pages)

    # Get all funds for dropdown
    funds = FundAccount.objects.filter(is_deleted=False, is_active=True).order_by("name")

    return render(
        request,
        "admin/allocation_rules.html",
        {
            "rules": rules_page,
            "page_obj": rules_page,
            "funds": funds,
        },
    )


@login_required
@admin_required
def add_allocation_rule_view(request):
    """Create a new allocation rule"""
    if request.method == "POST":
        try:
            trigger_event = request.POST.get("trigger_event")
            fund_id = request.POST.get("fund_id")
            allocation_type = request.POST.get("allocation_type")
            amount = request.POST.get("amount", "").strip()
            percentage = request.POST.get("percentage", "").strip()
            description = request.POST.get("description", "").strip()

            # Validate required fields
            if not trigger_event:
                return JsonResponse({"success": False, "error": "Trigger event is required"})
            if not fund_id:
                return JsonResponse({"success": False, "error": "Fund is required"})
            if not allocation_type:
                return JsonResponse({"success": False, "error": "Allocation type is required"})

            # Validate allocation amount/percentage
            if allocation_type == "fixed":
                if not amount:
                    return JsonResponse({"success": False, "error": "Amount is required for fixed allocation"})
                amount_decimal = Decimal(str(amount))
                if amount_decimal <= 0:
                    return JsonResponse({"success": False, "error": "Amount must be greater than zero"})
                percentage_decimal = None
            else:
                if not percentage:
                    return JsonResponse({"success": False, "error": "Percentage is required for percentage allocation"})
                percentage_decimal = Decimal(str(percentage))
                if percentage_decimal <= 0 or percentage_decimal > 100:
                    return JsonResponse({"success": False, "error": "Percentage must be between 0 and 100"})
                amount_decimal = None

            # Validate fund exists
            try:
                fund = FundAccount.objects.get(id=fund_id, is_deleted=False)
            except FundAccount.DoesNotExist:
                return JsonResponse({"success": False, "error": "Fund not found"})

            rule = FundAllocationRule.objects.create(
                trigger_event=trigger_event,
                fund=fund,
                allocation_type=allocation_type,
                amount=amount_decimal,
                percentage=percentage_decimal,
                description=description or None,
                created_by=request.user,
            )

            log_action(
                request,
                "create",
                "fund",
                rule.id,
                f"Created allocation rule: {rule.get_trigger_event_display()} → {fund.name}",
            )
            return JsonResponse({"success": True, "rule_id": rule.id})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def edit_allocation_rule_view(request, rule_id):
    """Update an allocation rule"""
    if request.method == "POST":
        try:
            rule = FundAllocationRule.objects.get(id=rule_id)

            trigger_event = request.POST.get("trigger_event")
            fund_id = request.POST.get("fund_id")
            allocation_type = request.POST.get("allocation_type")
            amount = request.POST.get("amount", "").strip()
            percentage = request.POST.get("percentage", "").strip()
            description = request.POST.get("description", "").strip()
            is_active = request.POST.get("is_active") == "true"

            # Validate required fields
            if not trigger_event:
                return JsonResponse({"success": False, "error": "Trigger event is required"})
            if not fund_id:
                return JsonResponse({"success": False, "error": "Fund is required"})
            if not allocation_type:
                return JsonResponse({"success": False, "error": "Allocation type is required"})

            # Validate allocation amount/percentage
            if allocation_type == "fixed":
                if not amount:
                    return JsonResponse({"success": False, "error": "Amount is required for fixed allocation"})
                amount_decimal = Decimal(str(amount))
                if amount_decimal <= 0:
                    return JsonResponse({"success": False, "error": "Amount must be greater than zero"})
                percentage_decimal = None
            else:
                if not percentage:
                    return JsonResponse({"success": False, "error": "Percentage is required for percentage allocation"})
                percentage_decimal = Decimal(str(percentage))
                if percentage_decimal <= 0 or percentage_decimal > 100:
                    return JsonResponse({"success": False, "error": "Percentage must be between 0 and 100"})
                amount_decimal = None

            # Validate fund exists
            try:
                fund = FundAccount.objects.get(id=fund_id, is_deleted=False)
            except FundAccount.DoesNotExist:
                return JsonResponse({"success": False, "error": "Fund not found"})

            rule.trigger_event = trigger_event
            rule.fund = fund
            rule.allocation_type = allocation_type
            rule.amount = amount_decimal
            rule.percentage = percentage_decimal
            rule.description = description or None
            rule.is_active = is_active
            rule.save()

            log_action(
                request,
                "update",
                "fund",
                rule.id,
                f"Updated allocation rule: {rule.get_trigger_event_display()} → {fund.name}",
            )
            return JsonResponse({"success": True})
        except FundAllocationRule.DoesNotExist:
            return JsonResponse({"success": False, "error": "Allocation rule not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


@login_required
@admin_required
def delete_allocation_rule_view(request, rule_id):
    """Delete an allocation rule"""
    if request.method == "POST":
        try:
            rule = FundAllocationRule.objects.get(id=rule_id)
            rule_desc = f"{rule.get_trigger_event_display()} → {rule.fund.name}"
            rule.delete()

            log_action(
                request,
                "delete",
                "fund",
                rule_id,
                f"Deleted allocation rule: {rule_desc}",
            )
            return JsonResponse({"success": True})
        except FundAllocationRule.DoesNotExist:
            return JsonResponse({"success": False, "error": "Allocation rule not found"})
        except Exception as e:
            return JsonResponse({"success": False, "error": str(e)})

    return JsonResponse({"success": False, "error": "Invalid request method"})


# ============================================================================
# Financial Periods (Phase A5)
# ----------------------------------------------------------------------------
# Promoted the orphan `templates/admin/financial_periods.html` to /finance/periods/.
# Business logic lives in `accounts.services.financial_period`; these views
# just translate POST forms into service calls and surface ServiceError
# messages as JSON for the inline-JS modal flow.
# ============================================================================

from django.utils.dateparse import parse_date as _parse_date_fp  # noqa: E402

from accounts.models import FinancialPeriod as _FinancialPeriod  # noqa: E402
from accounts.services import financial_period as _fp_service  # noqa: E402
from accounts.services.exceptions import ConflictError as _FPConflictError  # noqa: E402
from accounts.services.exceptions import NotFoundError as _FPNotFoundError  # noqa: E402
from accounts.services.exceptions import PermissionDeniedError as _FPPermissionDeniedError  # noqa: E402
from accounts.services.exceptions import ServiceError as _FPServiceError  # noqa: E402
from accounts.utils import _get_client_ip as _fp_client_ip  # noqa: E402


def _fp_service_error_status(exc):
    """Map a `ServiceError` subclass to an HTTP status code for JSON responses."""
    if isinstance(exc, _FPNotFoundError):
        return 404
    if isinstance(exc, _FPPermissionDeniedError):
        return 403
    if isinstance(exc, _FPConflictError):
        return 409
    return 400


@login_required
@admin_required
def financial_periods_view(request):
    """List page for financial periods at `/finance/periods/`."""
    periods = _FinancialPeriod.objects.all().order_by("-start_date")
    active_period = _fp_service.active()
    can_open_new = active_period is None or active_period.status == "closed"
    return render(
        request,
        "admin/financial_periods.html",
        {
            "periods": periods,
            "active_period": active_period,
            "can_open_new": can_open_new,
        },
    )


@login_required
@admin_required
def open_financial_period_view(request):
    """POST handler: open a new FY (calls `services.financial_period.open_period`)."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"}, status=405)
    name = (request.POST.get("name") or request.POST.get("label") or "").strip()
    start_date = _parse_date_fp((request.POST.get("start_date") or "").strip())
    end_date = _parse_date_fp((request.POST.get("end_date") or "").strip())
    try:
        period = _fp_service.open_period(
            name=name,
            start_date=start_date,
            end_date=end_date,
            actor=request.user,
            ip_address=_fp_client_ip(request),
            audit_via="panel",
        )
    except _FPServiceError as exc:
        return JsonResponse(
            {"success": False, "error": exc.message},
            status=_fp_service_error_status(exc),
        )
    return JsonResponse({"success": True, "period_id": period.id, "label": period.label, "status": period.status})


@login_required
@admin_required
def continue_financial_period_view(request, pk):
    """POST handler: re-activate a non-closed period."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"}, status=405)
    try:
        period = _fp_service.continue_period(
            period_id=pk,
            actor=request.user,
            ip_address=_fp_client_ip(request),
            audit_via="panel",
        )
    except _FPServiceError as exc:
        return JsonResponse(
            {"success": False, "error": exc.message},
            status=_fp_service_error_status(exc),
        )
    return JsonResponse({"success": True, "period_id": period.id, "label": period.label, "is_active": period.is_active})


@login_required
@admin_required
def close_financial_period_view(request, pk):
    """POST handler: close a period and lock its P&L snapshot."""
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Invalid request method"}, status=405)
    try:
        result = _fp_service.close_period(
            period_id=pk,
            actor=request.user,
            ip_address=_fp_client_ip(request),
            audit_via="panel",
        )
    except _FPServiceError as exc:
        return JsonResponse(
            {"success": False, "error": exc.message},
            status=_fp_service_error_status(exc),
        )
    period = result["period"]
    pnl = result["pnl"]
    society_account = result["society_account"]
    return JsonResponse(
        {
            "success": True,
            "period_id": period.id,
            "label": period.label,
            "status": period.status,
            "profit_and_loss_id": pnl.id if pnl else None,
            "society_account_id": society_account.id if society_account else None,
        }
    )
